
import io
import re
import zipfile
import sqlite3
from datetime import datetime
import json
import csv
import hashlib
import hmac
import base64
import time
import secrets
import random
import urllib.request
import urllib.parse
import urllib.error
from datetime import datetime as dt_datetime
from pathlib import Path
from datetime import date, datetime
from zoneinfo import ZoneInfo

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    BaseDocTemplate, Frame, PageTemplate, Table, TableStyle,
    Paragraph, Spacer, Image, PageBreak
)

BLACK = colors.black
RED = colors.HexColor("#E00000")
WHITE = colors.white

CATEGORIES = ["Cups", "Powders", "Torani", "Pastries", "Store"]

APP_DIR = Path(__file__).resolve().parent
BUILTIN_LOGO = APP_DIR / "apayao_brew_logo.png"
USERS_FILE = APP_DIR / "users.json"
ACTIVITY_LOG = APP_DIR / "activity_log.csv"

CATEGORY_TEMPLATES = {
    "Pastries": APP_DIR / "Pastries_Template.xlsx",
    "Cups": APP_DIR / "Cups_Template.xlsx",
    "Powders": APP_DIR / "Powders_Template.xlsx",
    "Torani": APP_DIR / "Torani_Template.xlsx",
    "Store": APP_DIR / "Store_Template.xlsx",
}
TEMPLATE_FILES = CATEGORY_TEMPLATES




BRANCH_CODES = {
    "RAMON": "RAM", "MUNOZ": "MUN", "MUÑOZ": "MUN",
    "SAN JOSE": "SJC", "BAYOMBONG": "BBO", "CORDON": "COR",
    "ECHAGUE": "ECH", "TABUK": "TAB", "BUNTUN": "BUN",
    "GATTARAN": "GAT", "STA. TERESITA": "STA", "STA TERESITA": "STA",
    "CAMALANIUGAN": "CAM", "BALLESTEROS": "BAL", "LUNA": "LUN",
    "SANCHEZ": "SAN", "GONZAGA": "GON", "NAMABBALAN": "NAM"
}


def hash_password(password, salt=None):
    if salt is None:
        salt = secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac(
        "sha256",
        password.encode("utf-8"),
        salt.encode("utf-8"),
        200_000
    ).hex()
    return salt, digest


def verify_password(password, salt, expected_hash):
    _, digest = hash_password(password, salt)
    return secrets.compare_digest(digest, expected_hash)


def save_users(users):
    USERS_FILE.write_text(json.dumps(users, indent=2), encoding="utf-8")


def load_users():
    if not USERS_FILE.exists():
        salt, digest = hash_password("ApayaoAdmin123!")
        users = {
            "admin": {
                "password_salt": salt,
                "password_hash": digest,
                "role": "Super Admin",
                "branch": "Main Store / All Branches",
                "active": True,
                "must_change_password": True
            }
        }
        save_users(users)
        return users

    try:
        return json.loads(USERS_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}


def log_activity(username, action, details=""):
    new_file = not ACTIVITY_LOG.exists()
    with ACTIVITY_LOG.open("a", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        if new_file:
            writer.writerow(["timestamp", "username", "action", "details"])
        writer.writerow([
            dt_datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            username,
            action,
            details
        ])



# ========================= PERMANENT USER ACCOUNTS (SUPABASE) =========================
def _auth_supabase_config():
    try:
        return str(st.secrets.get("SUPABASE_URL", "")).rstrip("/"), str(st.secrets.get("SUPABASE_SECRET_KEY", ""))
    except Exception:
        return "", ""

def _auth_sb_request(table, method="GET", params=None, payload=None):
    url, key = _auth_supabase_config()
    if not url or not key:
        raise RuntimeError("Supabase user storage is not configured.")
    endpoint = f"{url}/rest/v1/{table}"
    if params:
        endpoint += "?" + urllib.parse.urlencode(params, doseq=True)
    headers = {"apikey": key, "Authorization": f"Bearer {key}", "Content-Type": "application/json", "Accept": "application/json", "Prefer": "return=representation"}
    data = None if payload is None else json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(endpoint, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            body = resp.read().decode("utf-8")
            return json.loads(body) if body else []
    except urllib.error.HTTPError as e:
        detail = e.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Supabase error {e.code}: {detail}")

def supabase_user(username):
    try:
        rows = _auth_sb_request("app_users", "GET", {"select":"*", "username":f"eq.{username}", "limit":"1"})
        return rows[0] if rows else None
    except Exception:
        return None

def supabase_users():
    try:
        return _auth_sb_request("app_users", "GET", {"select":"*", "order":"username.asc"})
    except Exception:
        return []

def claimed_branches():
    return {str(r.get("branch") or "").strip() for r in supabase_users() if r.get("active", True) and str(r.get("branch") or "").strip() and r.get("role", "User") == "User"}

def create_supabase_user(username, full_name, branch, password, role="User", must_change=False):
    salt, digest = hash_password(password)
    payload = [{"username":username.strip(), "full_name":full_name.strip(), "branch":branch, "role":role, "password_salt":salt, "password_hash":digest, "active":True, "must_change_password":must_change, "updated_at":datetime.now().isoformat()}]
    return _auth_sb_request("app_users", "POST", payload=payload)

def update_supabase_user(username, values):
    values = dict(values); values["updated_at"] = datetime.now().isoformat()
    return _auth_sb_request("app_users", "PATCH", {"username":f"eq.{username}"}, values)

def activation_available_branches():
    """Load active branches without depending on functions defined later in the app."""
    branches = []
    try:
        db_path = APP_DIR / "apayao_v13.db"
        with sqlite3.connect(db_path) as c:
            rows = c.execute("SELECT name FROM branches_v13 WHERE active=1 ORDER BY name").fetchall()
            branches = [str(r[0]).strip() for r in rows if r and str(r[0]).strip()]
    except Exception:
        branches = []
    return sorted(dict.fromkeys(branches))

def activation_screen():
    st.subheader("Activate Branch Account")
    st.info("Select your assigned branch, enter your name, then create your personal login.")
    claimed = claimed_branches()
    branches = [b for b in activation_available_branches() if b not in claimed]
    if not branches:
        st.warning("All active branches already have an account. Please contact Super Admin.")
        if st.button("Back to Login"):
            st.session_state.pop("activation_mode", None); st.rerun()
        return
    branch = st.selectbox("Assigned Branch *", branches)
    full_name = st.text_input("Your Name *").strip()
    username = st.text_input("Create Personal Username *").strip()
    pw1 = st.text_input("Create New Password *", type="password")
    pw2 = st.text_input("Confirm New Password *", type="password")
    if st.button("ACTIVATE ACCOUNT", type="primary", use_container_width=True):
        if not full_name or not username:
            st.error("Name and username are required.")
        elif username.lower() in {"abrew", "admin"}:
            st.error("Please choose a different personal username.")
        elif len(pw1) < 8:
            st.error("Password must be at least 8 characters.")
        elif pw1 != pw2:
            st.error("Passwords do not match.")
        elif supabase_user(username):
            st.error("That username already exists.")
        elif branch in claimed_branches():
            st.error("That branch has already been claimed. Ask Super Admin to reset it if needed.")
        else:
            try:
                create_supabase_user(username, full_name, branch, pw1)
                st.success("Account activated. You can now log in using your personal username and password.")
                st.session_state.pop("activation_mode", None)
            except Exception as e:
                st.error(str(e))


def _persistent_login_secret():
    # Use the server-side Supabase secret as the signing key. It is never sent to the browser.
    _url, key = _auth_supabase_config()
    return key

def _make_login_token(username, record, hours=12):
    key = _persistent_login_secret()
    if not key:
        return ""
    exp = int(time.time()) + int(hours * 3600)
    fingerprint = str(record.get("password_hash", ""))[:16]
    payload = f"{username}|{exp}|{fingerprint}"
    sig = hmac.new(key.encode("utf-8"), payload.encode("utf-8"), hashlib.sha256).hexdigest()
    raw = f"{payload}|{sig}".encode("utf-8")
    return base64.urlsafe_b64encode(raw).decode("ascii").rstrip("=")

def _restore_persistent_login():
    if st.session_state.get("authenticated"):
        return True
    try:
        token = str(st.query_params.get("session", "") or "").strip()
    except Exception:
        token = ""
    if not token:
        return False
    try:
        padded = token + "=" * (-len(token) % 4)
        decoded = base64.urlsafe_b64decode(padded.encode("ascii")).decode("utf-8")
        username, exp_text, fingerprint, sig = decoded.rsplit("|", 3)
        if int(exp_text) < int(time.time()):
            st.query_params.pop("session", None)
            return False
        key = _persistent_login_secret()
        if not key:
            return False
        payload = f"{username}|{exp_text}|{fingerprint}"
        expected = hmac.new(key.encode("utf-8"), payload.encode("utf-8"), hashlib.sha256).hexdigest()
        if not hmac.compare_digest(sig, expected):
            st.query_params.pop("session", None)
            return False
        record = supabase_user(username) or load_users().get(username)
        if not record or not record.get("active", True):
            st.query_params.pop("session", None)
            return False
        if str(record.get("password_hash", ""))[:16] != fingerprint:
            st.query_params.pop("session", None)
            return False
        st.session_state["authenticated"] = True
        st.session_state["username"] = username
        st.session_state["role"] = record.get("role", "User")
        st.session_state["branch"] = record.get("branch", "")
        st.session_state["must_change_password"] = record.get("must_change_password", False)
        return True
    except Exception:
        try:
            st.query_params.pop("session", None)
        except Exception:
            pass
        return False

def login_screen():
    if st.session_state.get("activation_mode"):
        activation_screen(); return
    left, center, right = st.columns([1, 1.25, 1])
    with center:
        if BUILTIN_LOGO.exists(): st.image(str(BUILTIN_LOGO), width=115)
        st.markdown('<div class="login-brand"><div class="login-title">APAYAO BREW</div><div class="login-tagline">THE PERFECT BLEND.</div></div>', unsafe_allow_html=True)
    username = st.text_input("Username", placeholder="Enter your username")
    password = st.text_input("Password", type="password", placeholder="Enter your password")
    if st.button("Login", type="primary", use_container_width=True):
        if username == "abrew" and password == "abuser0123":
            st.session_state["activation_mode"] = True; st.rerun()
        record = supabase_user(username)
        if not record:
            record = load_users().get(username)
        if record and record.get("active", True) and verify_password(password, record["password_salt"], record["password_hash"]):
            st.session_state["authenticated"] = True
            st.session_state["username"] = username
            st.session_state["role"] = record.get("role", "User")
            st.session_state["branch"] = record.get("branch", "")
            st.session_state["must_change_password"] = record.get("must_change_password", False)
            token = _make_login_token(username, record, hours=12)
            if token:
                st.query_params["session"] = token
            log_activity(username, "LOGIN"); st.rerun()
        else:
            st.error("Invalid username or password.")

def change_password_form(username, forced=False):
    st.subheader("Change Password")
    if forced:
        st.warning("You must change the default password before using the app.")

    new_pw = st.text_input("New Password", type="password", key="new_pw")
    confirm_pw = st.text_input("Confirm New Password", type="password", key="confirm_pw")

    if st.button("Save New Password", type="primary"):
        if len(new_pw) < 8:
            st.error("Password must be at least 8 characters."); return
        if new_pw != confirm_pw:
            st.error("Passwords do not match."); return
        salt, digest = hash_password(new_pw)
        # Permanent accounts live in Supabase so password changes survive redeploys.
        if supabase_user(username):
            update_supabase_user(username, {"password_salt":salt, "password_hash":digest, "must_change_password":False})
        else:
            users = load_users()
            if username not in users:
                st.error("User account not found."); return
            users[username]["password_salt"] = salt
            users[username]["password_hash"] = digest
            users[username]["must_change_password"] = False
            save_users(users)
        st.session_state["must_change_password"] = False
        log_activity(username, "PASSWORD_CHANGED")
        st.success("Password changed successfully.")
        st.rerun()


def admin_panel():
    st.subheader("Admin — User Management")
    st.caption("Accounts created here are stored permanently in Supabase and survive app refreshes/redeploys.")

    with st.expander("Create New User", expanded=False):
        new_user = st.text_input("New Username", key="admin_new_username")
        full_name = st.text_input("Name", key="admin_new_full_name")
        new_role = st.selectbox("Role", ["User", "Admin", "Super Admin"], key="admin_new_role")
        branch_options = ["Main Store / All Branches"] + available_branches()
        default_branch = "Main Store / All Branches" if new_role in ("Admin", "Super Admin") else branch_options[0]
        new_branch = st.selectbox("Branch", branch_options, index=branch_options.index(default_branch), key="admin_new_branch")
        temp_pw = st.text_input("Temporary Password", type="password", key="admin_temp_pw")
        if st.button("Create User"):
            clean_user = new_user.strip()
            if not clean_user:
                st.error("Username is required.")
            elif supabase_user(clean_user) or clean_user in load_users():
                st.error("Username already exists.")
            elif len(temp_pw) < 8:
                st.error("Temporary password must be at least 8 characters.")
            else:
                try:
                    create_supabase_user(clean_user, full_name.strip() or clean_user, new_branch, temp_pw, role=new_role, must_change=True)
                    log_activity(st.session_state["username"], "USER_CREATED", f"{clean_user} ({new_role})")
                    st.success(f"User '{clean_user}' created permanently.")
                    st.rerun()
                except Exception as e:
                    st.error(str(e))

    permanent = supabase_users()
    if permanent:
        st.markdown("### Permanent Accounts")
        st.dataframe(pd.DataFrame([{"Username":r.get("username"),"Name":r.get("full_name"),"Branch":r.get("branch"),"Role":r.get("role"),"Active":r.get("active",True),"Must Change Password":r.get("must_change_password",False)} for r in permanent]), use_container_width=True, hide_index=True)
        target = st.selectbox("Select Permanent User to Manage", [r.get("username") for r in permanent], key="perm_manage")
        selected = next(r for r in permanent if r.get("username") == target)
        c1,c2,c3 = st.columns(3)
        with c1:
            if st.button("Toggle Active/Inactive"):
                if target == st.session_state["username"]:
                    st.error("You cannot deactivate your own account.")
                else:
                    update_supabase_user(target,{"active":not selected.get("active",True)})
                    log_activity(st.session_state["username"],"USER_STATUS_CHANGED",f"{target}: active={not selected.get('active',True)}")
                    st.rerun()
        with c2:
            reset_pw = st.text_input("Reset Password",type="password",key=f"perm_reset_{target}",placeholder="New temporary password")
            if st.button("Reset Selected User Password"):
                if len(reset_pw)<8:
                    st.error("Password must be at least 8 characters.")
                else:
                    salt,digest=hash_password(reset_pw)
                    update_supabase_user(target,{"password_salt":salt,"password_hash":digest,"must_change_password":True})
                    log_activity(st.session_state["username"],"PASSWORD_RESET",target)
                    st.success("Password reset. User must change it on next login.")
        with c3:
            roles=["User","Admin","Super Admin"]
            role_now=selected.get("role","User") if selected.get("role","User") in roles else "User"
            role_new=st.selectbox("Change Role",roles,index=roles.index(role_now),key=f"perm_role_{target}")
            if st.button("Update Role"):
                if target == st.session_state["username"] and role_new != selected.get("role"):
                    st.error("You cannot change your own role while logged in.")
                else:
                    update_supabase_user(target,{"role":role_new})
                    log_activity(st.session_state["username"],"ROLE_CHANGED",f"{target}: {role_new}")
                    st.rerun()
        branch_options=["Main Store / All Branches"]+available_branches()
        current_branch=selected.get("branch") or "Main Store / All Branches"
        if current_branch not in branch_options: branch_options.append(current_branch)
        assigned=st.selectbox("Assigned Branch",branch_options,index=branch_options.index(current_branch),key=f"perm_branch_{target}")
        if st.button("Update Assigned Branch"):
            update_supabase_user(target,{"branch":assigned})
            log_activity(st.session_state["username"],"BRANCH_CHANGED",f"{target}: {assigned}")
            st.success("Assigned branch updated."); st.rerun()
    else:
        st.info("No permanent accounts found yet.")

    legacy=load_users()
    if legacy:
        st.divider()
        st.markdown("### Legacy / Recovery Account")
        st.caption("Kept as a recovery fallback. New accounts are no longer stored here.")
        st.dataframe(pd.DataFrame([{"Username":u,"Role":r.get("role","User"),"Branch":r.get("branch",""),"Active":r.get("active",True)} for u,r in legacy.items()]),use_container_width=True,hide_index=True)

    st.subheader("Activity Log")
    if ACTIVITY_LOG.exists():
        try:
            log_df=pd.read_csv(ACTIVITY_LOG)
            st.dataframe(log_df.sort_index(ascending=False).head(200),use_container_width=True,hide_index=True)
        except Exception:
            st.info("Activity log is not available yet.")
    else:
        st.info("No activity recorded yet.")


st.set_page_config(
    page_title="Apayao Brew Delivery Receipt Generator",
    page_icon="☕",
    layout="wide"
)

st.markdown("""
<style>
:root {
    --ab-brown: #5A3F2D;
    --ab-brown-dark: #3C271B;
    --ab-brown-soft: #765843;
    --ab-cream: #FBF7F2;
    --ab-beige: #EEE3D8;
    --ab-line: #E4D8CD;
    --ab-white: #FFFFFF;
    --ab-text: #33261F;
}

/* Entire app */
.stApp {
    background: linear-gradient(180deg, #FCF9F6 0%, #FFFFFF 52%, #FAF5F0 100%);
    color: var(--ab-text);
}
.block-container {
    max-width: 1220px;
    padding-top: 1.0rem;
    padding-bottom: 3rem;
}
h1, h2, h3 {
    color: var(--ab-brown-dark) !important;
    font-family: Georgia, "Times New Roman", serif;
}
p, label, .stMarkdown {
    color: var(--ab-text);
}

/* Sidebar */
section[data-testid="stSidebar"] {
    background: #FFFCF8;
    border-right: 1px solid var(--ab-line);
}
section[data-testid="stSidebar"] .block-container {
    padding-top: 1rem;
}
section[data-testid="stSidebar"] button {
    border-radius: 10px !important;
}

/* Buttons */
.stButton > button,
.stDownloadButton > button {
    border-radius: 12px !important;
    border: 1px solid #D8C8BB !important;
    min-height: 45px;
    font-weight: 650 !important;
    transition: all .15s ease-in-out;
}
.stButton > button:hover,
.stDownloadButton > button:hover {
    border-color: var(--ab-brown) !important;
    box-shadow: 0 6px 18px rgba(90, 63, 45, .12);
    transform: translateY(-1px);
}
button[kind="primary"] {
    background: var(--ab-brown) !important;
    color: white !important;
    border-color: var(--ab-brown) !important;
}

/* Inputs */
div[data-baseweb="input"] > div,
div[data-baseweb="select"] > div,
div[data-baseweb="base-input"] {
    border-radius: 10px !important;
}
div[data-testid="stFileUploader"] {
    background: white;
    border: 1px dashed #BFA998;
    padding: 12px;
    border-radius: 14px;
}

/* Tables */
div[data-testid="stDataFrame"] {
    border: 1px solid var(--ab-line);
    border-radius: 14px;
    overflow: hidden;
    background: white;
}

/* Alerts */
div[data-testid="stAlert"] {
    border-radius: 12px;
}

/* Brand/header */
.ab-brandbar {
    display:flex;
    align-items:center;
    justify-content:space-between;
    gap:16px;
    padding: 9px 4px 13px 4px;
    border-bottom: 1px solid var(--ab-line);
    margin-bottom: 12px;
}
.ab-brandname {
    font-weight:800;
    letter-spacing:.18em;
    font-size:1.05rem;
    color:var(--ab-brown-dark);
}
.ab-brandtag {
    font-size:.72rem;
    letter-spacing:.28em;
    color:var(--ab-brown-soft);
    margin-top:2px;
}
.ab-user {
    text-align:right;
    color:#6B5547;
    font-size:.84rem;
}

/* Hero */
.ab-hero {
    border-radius: 18px;
    padding: 28px 32px;
    margin: 8px 0 24px 0;
    background:
        radial-gradient(circle at 88% 25%, rgba(255,255,255,.28) 0 9%, transparent 10%),
        linear-gradient(110deg, #EEE0D2 0%, #F8F1EA 44%, #C6A98F 100%);
    border: 1px solid #DDCDBF;
    box-shadow: 0 12px 30px rgba(72, 48, 32, .08);
}
.ab-welcome {
    font-family: Georgia, serif;
    font-style: italic;
    color: #6B4A35;
    font-size: 1.05rem;
}
.ab-hero h1 {
    font-size: 2.6rem !important;
    line-height: 1.05;
    margin: 3px 0 8px 0;
}
.ab-hero p {
    font-size: 1rem;
    color: #644C3D;
    margin: 0 0 13px 0;
}
.ab-perfect {
    letter-spacing: .32em;
    font-weight: 800;
    color: var(--ab-brown-dark);
    font-size: .78rem;
}

/* Step cards */
.ab-sectiontitle {
    font-family: Georgia, serif;
    font-weight:800;
    font-size:1.55rem;
    color:var(--ab-brown-dark);
    margin-top: 18px;
}
.ab-sectionhint {
    color:#776255;
    margin-bottom:12px;
}
.ab-steps {
    display:grid;
    grid-template-columns: repeat(4, 1fr);
    gap:12px;
    margin: 10px 0 22px 0;
}
.ab-step {
    background:#fff;
    border:1px solid var(--ab-line);
    border-radius:14px;
    padding:15px 14px;
    min-height:115px;
    box-shadow:0 5px 15px rgba(72,48,32,.04);
}
.ab-stepnum {
    background:var(--ab-brown);
    color:#fff;
    height:28px;
    width:28px;
    display:flex;
    align-items:center;
    justify-content:center;
    border-radius:50%;
    font-weight:700;
    margin-bottom:10px;
}
.ab-step b {
    color:var(--ab-brown-dark);
}
.ab-step span {
    display:block;
    color:#756357;
    font-size:.85rem;
    margin-top:5px;
}

/* Category radio styled like cards */
div[role="radiogroup"] {
    display:flex !important;
    gap:10px !important;
    flex-wrap:wrap !important;
}
div[role="radiogroup"] > label {
    background:#fff !important;
    border:1px solid var(--ab-line) !important;
    border-radius:14px !important;
    padding:14px 18px !important;
    min-width:135px;
    min-height:58px;
    box-shadow:0 4px 13px rgba(72,48,32,.04);
}
div[role="radiogroup"] > label:hover {
    border-color:#A78973 !important;
}
div[role="radiogroup"] > label:has(input:checked) {
    background:var(--ab-brown) !important;
    color:white !important;
    border-color:var(--ab-brown) !important;
}
div[role="radiogroup"] > label:has(input:checked) p {
    color:white !important;
    font-weight:700 !important;
}

/* Login */
.login-brand {
    text-align:center;
    margin-bottom:18px;
}
.login-title {
    font-size:1.55rem;
    letter-spacing:.16em;
    font-weight:900;
    color:var(--ab-brown-dark);
}
.login-tagline {
    letter-spacing:.28em;
    font-size:.72rem;
    font-weight:700;
    color:var(--ab-brown-soft);
    margin-top:3px;
}
.login-subtitle {
    font-family:Georgia, serif;
    font-size:1.18rem;
    font-weight:700;
    color:var(--ab-brown-dark);
    margin-top:16px;
}

/* Footer */
.ab-footer {
    border-top:1px solid var(--ab-line);
    margin-top:34px;
    padding:18px 0 0 0;
    display:flex;
    justify-content:space-between;
    color:#7A6659;
    font-size:.76rem;
    letter-spacing:.04em;
}

@media (max-width: 900px) {
    .ab-steps { grid-template-columns:1fr 1fr; }
    .ab-hero h1 { font-size:2rem !important; }
}
</style>
""", unsafe_allow_html=True)

# Authentication gate. Restore a signed 12-hour login session after normal browser refresh/reopen.
if not st.session_state.get("authenticated"):
    _restore_persistent_login()
if not st.session_state.get("authenticated"):
    login_screen()
    st.stop()

current_user = st.session_state["username"]
current_role = st.session_state["role"]

if st.session_state.get("must_change_password"):
    change_password_form(current_user, forced=True)
    st.stop()

with st.sidebar:
    if BUILTIN_LOGO.exists():
        st.image(str(BUILTIN_LOGO), width=82)
    st.markdown("### APAYAO BREW")
    st.caption("THE PERFECT BLEND.")
    st.divider()
    st.write(f"**Logged in as:** {current_user}")
    st.write(f"**Role:** {current_role}")
    if st.button("Change My Password"):
        st.session_state["show_change_password"] = not st.session_state.get("show_change_password", False)
    if st.button("Logout"):
        log_activity(current_user, "LOGOUT")
        for key in ["authenticated", "username", "role", "branch", "must_change_password"]:
            st.session_state.pop(key, None)
        try:
            st.query_params.pop("session", None)
        except Exception:
            pass
        st.rerun()

if st.session_state.get("show_change_password"):
    change_password_form(current_user, forced=False)


def txt(v):
    if v is None or pd.isna(v):
        return ""
    return str(v).strip()


def safe_filename(v):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", str(v)).strip("_") or "delivery_receipt"


def is_number(v):
    try:
        float(v)
        return True
    except Exception:
        return False


def excel_date(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            return (pd.Timestamp("1899-12-30") + pd.to_timedelta(float(v), unit="D")).date()
        except Exception:
            pass
    try:
        return pd.to_datetime(v).date()
    except Exception:
        return None


def find_label_value(raw, labels):
    normalized = {x.upper().rstrip(":") for x in labels}
    for r in range(len(raw)):
        for c in range(len(raw.columns)):
            val = txt(raw.iat[r, c]).upper().rstrip(":")
            if val in normalized:
                for cc in range(c + 1, min(c + 4, len(raw.columns))):
                    candidate = raw.iat[r, cc]
                    if txt(candidate):
                        return candidate
    return ""


def detect_matrix_pastries(raw):
    best = None
    best_score = -1

    for r in range(len(raw)):
        for branch_col in range(min(4, len(raw.columns) - 1)):
            left = txt(raw.iat[r, branch_col]).upper()
            if left not in {"", "BRANCH", "BRANCHES"}:
                continue

            product_names = []
            for c in range(branch_col + 1, len(raw.columns)):
                v = txt(raw.iat[r, c])
                if v and not is_number(v):
                    product_names.append(v)

            if len(product_names) < 2:
                continue

            branch_rows = 0
            numeric_cells = 0
            for rr in range(r + 1, min(r + 30, len(raw))):
                branch = txt(raw.iat[rr, branch_col])
                if not branch or branch.upper() in {"TOTAL", "TOTALS", "GRAND TOTAL"}:
                    continue
                nums = 0
                for cc in range(branch_col + 1, len(raw.columns)):
                    v = raw.iat[rr, cc]
                    if txt(v) and is_number(v):
                        nums += 1
                if nums:
                    branch_rows += 1
                    numeric_cells += nums

            score = len(product_names) * 5 + branch_rows * 4 + numeric_cells
            if branch_rows and score > best_score:
                best_score = score
                best = (r, branch_col)

    if not best:
        raise ValueError(
            "I could not find the branch-by-product table. "
            "Use the simple Excel format with branch names in rows and products in columns."
        )

    header_row, branch_col = best
    products = []

    for c in range(branch_col + 1, len(raw.columns)):
        name = txt(raw.iat[header_row, c])
        if name:
            products.append((c, name))

    branches = []
    for r in range(header_row + 1, len(raw)):
        branch = txt(raw.iat[r, branch_col])
        if not branch or branch.upper() in {"TOTAL", "TOTALS", "GRAND TOTAL"}:
            continue

        items = []
        numeric_seen = False
        for c, product in products:
            value = raw.iat[r, c]
            qty = 0.0
            if txt(value):
                try:
                    qty = float(value)
                    numeric_seen = True
                except Exception:
                    qty = 0.0
            items.append((product, qty))

        if numeric_seen:
            branches.append({"branch": branch, "items": items})

    if not branches:
        raise ValueError("No branch quantities were found.")
    return branches



def detect_matrix_transposed(raw):
    """
    Format for Cups, Powders, Torani, Store:
    - Branch names are across the uppermost/header row.
    - Item names are down the first column.
    - Quantities are in the grid.
    """

    best = None
    best_score = -1

    # Find a likely branch header row and item-name column.
    for r in range(len(raw)):
        for item_col in range(min(4, len(raw.columns) - 1)):
            left = txt(raw.iat[r, item_col]).upper()

            # Top-left can be blank or a label like ITEM / ITEMS.
            if left not in {"", "ITEM", "ITEMS", "ITEM DESCRIPTION", "PRODUCT", "PRODUCTS"}:
                continue

            branches = []
            for c in range(item_col + 1, len(raw.columns)):
                value = txt(raw.iat[r, c])
                if value and not is_number(value):
                    branches.append((c, value))

            if len(branches) < 2:
                continue

            item_rows = 0
            numeric_cells = 0

            for rr in range(r + 1, len(raw)):
                item = txt(raw.iat[rr, item_col])
                if not item:
                    continue
                if item.upper() in {"TOTAL", "TOTALS", "GRAND TOTAL"}:
                    continue

                nums = 0
                for c, _branch in branches:
                    v = raw.iat[rr, c]
                    if txt(v) and is_number(v):
                        nums += 1

                if nums:
                    item_rows += 1
                    numeric_cells += nums

            score = len(branches) * 5 + item_rows * 4 + numeric_cells
            if item_rows and score > best_score:
                best_score = score
                best = (r, item_col, branches)

    if not best:
        raise ValueError(
            "I could not find the category table. "
            "For Cups, Powders, Torani, and Store, branches must be across the top row "
            "and item names must be down the first column."
        )

    header_row, item_col, branch_headers = best

    # Create one branch record per branch column.
    result = []
    for branch_col, branch_name in branch_headers:
        items = []
        numeric_seen = False

        for r in range(header_row + 1, len(raw)):
            item = txt(raw.iat[r, item_col])
            if not item:
                continue
            if item.upper() in {"TOTAL", "TOTALS", "GRAND TOTAL"}:
                continue

            value = raw.iat[r, branch_col]
            qty = 0.0

            if txt(value):
                try:
                    qty = float(value)
                    numeric_seen = True
                except Exception:
                    qty = 0.0

            items.append((item, qty))

        if numeric_seen:
            result.append({
                "branch": branch_name,
                "items": items
            })

    if not result:
        raise ValueError("No branch quantities were found in the uploaded file.")

    return result


def detect_matrix_by_category(raw, category):
    if category == "Pastries":
        return detect_matrix_pastries(raw)
    return detect_matrix_transposed(raw)

def code_for_branch(branch):
    key = re.sub(r"\s+", " ", branch.upper().strip())
    if key in BRANCH_CODES:
        return BRANCH_CODES[key]
    letters = re.sub(r"[^A-Z]", "", key)
    return (letters[:3] or "BR").upper()


def format_date(d):
    if not d:
        return ""
    if isinstance(d, datetime):
        d = d.date()
    if isinstance(d, date):
        return d.strftime("%m/%d/%Y")
    return txt(d)


def dr_no_text(base_dr, delivery_date, branch):
    if not delivery_date:
        delivery_date = date.today()
    base = txt(base_dr)
    mmdd = delivery_date.strftime("%m%d")
    code = code_for_branch(branch)
    return f"DR #{base} {mmdd} {code}" if base else f"DR #{mmdd} {code}"


def qty_text(q):
    try:
        q = float(q)
        if q == 0:
            return ""
        return str(int(q)) if q.is_integer() else f"{q:g}"
    except Exception:
        return txt(q)


def infer_row_category(selected_category, product):
    p = product.upper()
    if selected_category == "Cups":
        if "LID" in p:
            return "LIDS"
        if "CUP" in p:
            return "CUPS"
    return selected_category.upper()


def make_styles():
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "title", parent=styles["Normal"], fontName="Helvetica-Bold",
            fontSize=11.5, leading=12.5, alignment=TA_CENTER, textColor=BLACK
        ),
        "normal": ParagraphStyle(
            "normal", parent=styles["Normal"], fontName="Helvetica",
            fontSize=7.4, leading=8.5, textColor=BLACK
        ),
        "small": ParagraphStyle(
            "small", parent=styles["Normal"], fontName="Helvetica",
            fontSize=6.2, leading=7.2, textColor=BLACK
        ),
        "center": ParagraphStyle(
            "center", parent=styles["Normal"], fontName="Helvetica",
            fontSize=7.2, leading=8.2, alignment=TA_CENTER, textColor=BLACK
        ),
        "boldcenter": ParagraphStyle(
            "boldcenter", parent=styles["Normal"], fontName="Helvetica-Bold",
            fontSize=7.2, leading=8.2, alignment=TA_CENTER, textColor=BLACK
        ),
        "dr": ParagraphStyle(
            "dr", parent=styles["Normal"], fontName="Helvetica-Bold",
            fontSize=14, leading=15, alignment=TA_CENTER, textColor=RED
        ),
        "note": ParagraphStyle(
            "note", parent=styles["Normal"], fontName="Helvetica-BoldOblique",
            fontSize=7.7, leading=8.5, textColor=RED
        ),
    }


def add_receipt_page(story, styles, branch, items, selected_category,
                     base_dr, store_origin, prepared_by, delivered_by,
                     delivery_date, approved_by, generated_by, logo_bytes=None):

    try:
        if BUILTIN_LOGO.exists():
            logo = Image(str(BUILTIN_LOGO), width=10*mm, height=10*mm)
            logo.hAlign = "CENTER"
            story.append(logo)
    except Exception:
        pass

    story.append(Paragraph("APAYAO BREW OPC", styles["title"]))
    story.append(Paragraph("DELIVERY FORM", styles["title"]))
    story.append(Spacer(1, 1.2*mm))

    # Branch and DR row, matching the user's layout.
    branch_block = Table([
        [
            Paragraph("<b>BRANCH:</b>", styles["normal"]),
            Paragraph(f"<i>{branch}</i>", styles["center"]),
            Paragraph(dr_no_text(base_dr, delivery_date, branch), styles["dr"])
        ]
    ], colWidths=[25*mm, 91*mm, 78*mm])
    branch_block.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("LINEBELOW",(1,0),(1,0),0.55,BLACK),
        ("LEFTPADDING",(0,0),(-1,-1),0),
        ("RIGHTPADDING",(0,0),(-1,-1),0),
        ("TOPPADDING",(0,0),(-1,-1),0),
        ("BOTTOMPADDING",(0,0),(-1,-1),0),
    ]))
    story.append(branch_block)
    story.append(Spacer(1, 3*mm))

    info = Table([
        [Paragraph("Store origin:", styles["normal"]),
         Paragraph(f"<b>{store_origin}</b>", styles["normal"]),
         "", "",
         "", ""],
        ["", "", "", "",
         Paragraph("Delivery Date:", styles["normal"]),
         Paragraph(format_date(delivery_date), styles["normal"])],
        [Paragraph("Prepared by:", styles["normal"]),
         Paragraph(f"<i>{prepared_by}</i>", styles["normal"]),
         "", "",
         Paragraph("Receiving Date/Shift:", styles["small"]),
         ""],
        [Paragraph("Delivery by:", styles["normal"]),
         Paragraph(f"<i>{delivered_by}</i>", styles["normal"]),
         "", "", "", ""],
    ], colWidths=[23*mm, 40*mm, 12*mm, 34*mm, 31*mm, 54*mm])
    info.setStyle(TableStyle([
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("LINEBELOW",(1,2),(1,3),0.45,BLACK),
        ("LINEBELOW",(5,1),(5,2),0.45,BLACK),
        ("LEFTPADDING",(0,0),(-1,-1),0),
        ("RIGHTPADDING",(0,0),(-1,-1),1),
        ("TOPPADDING",(0,0),(-1,-1),0.5),
        ("BOTTOMPADDING",(0,0),(-1,-1),0.5),
    ]))
    story.append(info)
    story.append(Spacer(1, 3*mm))

    headers = [
        Paragraph("<b>CATEGORY</b>", styles["boldcenter"]),
        Paragraph("<b>ORDER (qty)</b>", styles["small"]),
        Paragraph("<b>DELIVERY (qty)</b>", styles["small"]),
        Paragraph("<b>ITEM DESCRIPTION</b> <font size='5.5'>(Indicate brand name, product name, size,weight, and flavor)</font>", styles["center"]),
        Paragraph("<b>ACTUAL RECEIVED (qty)</b>", styles["small"]),
        Paragraph("<b>VARIANCE</b>", styles["boldcenter"]),
    ]

    rows = [headers]

    # Show only products with a quantity greater than zero.
    # Blank or zero quantities are omitted from the generated DR.
    active_items = []
    for product, qty in items:
        try:
            numeric_qty = float(qty)
        except Exception:
            numeric_qty = 0.0
        if numeric_qty > 0:
            active_items.append((product, numeric_qty))

    # Category appears only once, on the uppermost populated item row.
    for i, (product, qty) in enumerate(active_items):
        rows.append([
            Paragraph(selected_category.upper() if i == 0 else "", styles["normal"]),
            "",
            Paragraph(qty_text(qty), styles["center"]),
            Paragraph(product, styles["normal"]),
            "",
            ""
        ])

    # Keep the long receiving table even when some products are omitted.
    target_body_rows = 27
    for _ in range(max(0, target_body_rows - len(active_items))):
        rows.append(["", "", "", "", "", ""])

    t = Table(
        rows,
        colWidths=[24*mm, 20*mm, 22*mm, 80*mm, 28*mm, 24*mm],
        repeatRows=1
    )
    t.setStyle(TableStyle([
        ("GRID",(0,0),(-1,-1),0.55,BLACK),
        ("VALIGN",(0,0),(-1,-1),"MIDDLE"),
        ("ALIGN",(1,1),(2,-1),"CENTER"),
        ("ALIGN",(4,1),(5,-1),"CENTER"),
        ("LEFTPADDING",(0,0),(-1,-1),1.5),
        ("RIGHTPADDING",(0,0),(-1,-1),1.5),
        ("TOPPADDING",(0,0),(-1,0),2.2),
        ("BOTTOMPADDING",(0,0),(-1,0),2.2),
        ("TOPPADDING",(0,1),(-1,-1),2.0),
        ("BOTTOMPADDING",(0,1),(-1,-1),2.0),
    ]))
    story.append(t)

    times = Table([
        [
            Paragraph("TIME OF ARRIVAL:", styles["normal"]), "",
            Paragraph("TIME OF DEPARTURE:", styles["normal"]), ""
        ]
    ], colWidths=[34*mm, 62*mm, 37*mm, 65*mm])
    times.setStyle(TableStyle([
        ("LINEBELOW",(1,0),(1,0),0.5,BLACK),
        ("LINEBELOW",(3,0),(3,0),0.5,BLACK),
        ("LEFTPADDING",(0,0),(-1,-1),0),
        ("RIGHTPADDING",(0,0),(-1,-1),0),
        ("TOPPADDING",(0,0),(-1,-1),1),
        ("BOTTOMPADDING",(0,0),(-1,-1),1),
    ]))
    story.append(times)
    story.append(Spacer(1, 1.5*mm))

    story.append(Paragraph(
        "Note: On remarks column indicate discrepancy based on delivery versus actual received.",
        styles["note"]
    ))
    story.append(Spacer(1, 1.5*mm))
    story.append(Paragraph(
        f"Generated by: <b>{generated_by}</b>",
        styles["small"]
    ))
    story.append(Spacer(1, 2.5*mm))

    sign = Table([
        [Paragraph("Received by:", styles["normal"]), "", "", "",
         Paragraph("Supervisor on duty:", styles["normal"]), ""],
        ["", "", "", "", "", ""],
        [Paragraph("Cashier on duty", styles["center"]), "", "", "",
         Paragraph("(Fullname w/signature)", styles["small"]), ""],
        [Paragraph("(Fullname w/signature, Date)", styles["small"]), "", "", "", "", ""],
        ["", "", "", "", "", ""],
        ["", "", "", "", Paragraph("Approved by:", styles["normal"]), ""],
        ["", "", "", "", "", ""],
        ["", "", "", "", Paragraph(approved_by, styles["center"]), ""],
    ], colWidths=[48*mm, 18*mm, 23*mm, 20*mm, 45*mm, 44*mm])

    sign.setStyle(TableStyle([
        ("LINEBELOW",(0,1),(1,1),0.5,BLACK),
        ("LINEBELOW",(4,1),(5,1),0.5,BLACK),
        ("LINEBELOW",(4,6),(5,6),0.5,BLACK),
        ("LEFTPADDING",(0,0),(-1,-1),0),
        ("RIGHTPADDING",(0,0),(-1,-1),0),
        ("TOPPADDING",(0,0),(-1,-1),0.5),
        ("BOTTOMPADDING",(0,0),(-1,-1),0.5),
    ]))
    story.append(sign)


def build_combined_pdf(branches, selected_category, base_dr, store_origin,
                       prepared_by, delivered_by, delivery_date,
                       approved_by, generated_by, logo_bytes=None):
    buf = io.BytesIO()
    doc = BaseDocTemplate(
        buf, pagesize=A4,
        leftMargin=7*mm, rightMargin=7*mm,
        topMargin=6*mm, bottomMargin=6*mm
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="receipt", frames=frame)])

    styles = make_styles()
    story = []

    for i, b in enumerate(branches):
        add_receipt_page(
            story, styles,
            b["branch"], b["items"], selected_category,
            base_dr, store_origin, prepared_by, delivered_by,
            delivery_date, approved_by, generated_by, logo_bytes
        )
        if i < len(branches)-1:
            story.append(PageBreak())

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


def build_single_pdf(branch, selected_category, base_dr, store_origin,
                     prepared_by, delivered_by, delivery_date,
                     approved_by, generated_by, logo_bytes=None):
    return build_combined_pdf(
        [branch], selected_category, base_dr, store_origin,
        prepared_by, delivered_by, delivery_date,
        approved_by, generated_by, logo_bytes
    )



# ========================= V13.1 ORDER REQUEST / DELIVERY / RECEIVING =========================
DB_FILE = APP_DIR / "apayao_v13.db"
ORDER_CATEGORIES = ["Cups", "Powders", "Torani", "Pastries", "Store"]
DEFAULT_BRANCHES = ["Ramon", "Cordon", "Gattaran", "Camalaniugan", "Buntun", "Namabbalan"]

def db_conn():
    return sqlite3.connect(DB_FILE)

def init_v13_db():
    with db_conn() as c:
        c.executescript("""
        CREATE TABLE IF NOT EXISTS order_cycles(
          id INTEGER PRIMARY KEY AUTOINCREMENT, order_no TEXT UNIQUE NOT NULL,
          category TEXT NOT NULL, status TEXT DEFAULT 'Open', created_by TEXT,
          created_at TEXT, deadline TEXT, notes TEXT);
        CREATE TABLE IF NOT EXISTS order_cycle_branches(
          id INTEGER PRIMARY KEY AUTOINCREMENT, cycle_id INTEGER NOT NULL,
          branch TEXT NOT NULL, status TEXT DEFAULT 'Pending', submitted_by TEXT,
          submitted_at TEXT, UNIQUE(cycle_id,branch));
        CREATE TABLE IF NOT EXISTS branch_order_items(
          id INTEGER PRIMARY KEY AUTOINCREMENT, cycle_branch_id INTEGER NOT NULL,
          item TEXT NOT NULL, ordered REAL DEFAULT 0, allocated REAL DEFAULT 0,
          delivered REAL DEFAULT 0, received REAL DEFAULT 0,
          UNIQUE(cycle_branch_id,item));
        CREATE TABLE IF NOT EXISTS deliveries_v2(
          id INTEGER PRIMARY KEY AUTOINCREMENT, delivery_no TEXT UNIQUE,
          cycle_branch_id INTEGER NOT NULL, status TEXT DEFAULT 'For Receiving',
          created_by TEXT, created_at TEXT, received_by TEXT, received_at TEXT,
          remarks TEXT, proof_name TEXT);
        CREATE TABLE IF NOT EXISTS audit_v13(
          id INTEGER PRIMARY KEY AUTOINCREMENT, username TEXT, action TEXT,
          details TEXT, created_at TEXT);
        CREATE TABLE IF NOT EXISTS branches_v13(
          id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT UNIQUE NOT NULL,
          active INTEGER DEFAULT 1, created_by TEXT, created_at TEXT);
        CREATE TABLE IF NOT EXISTS emergency_orders(
          id INTEGER PRIMARY KEY AUTOINCREMENT, cycle_branch_id INTEGER NOT NULL,
          item TEXT NOT NULL, requested_qty REAL NOT NULL DEFAULT 0,
          approved_qty REAL, status TEXT DEFAULT 'Pending Approval', reason TEXT NOT NULL,
          requested_by TEXT, requested_at TEXT, reviewed_by TEXT, reviewed_at TEXT, review_note TEXT);
        CREATE TABLE IF NOT EXISTS order_branch_changes(
          id INTEGER PRIMARY KEY AUTOINCREMENT, cycle_id INTEGER NOT NULL, branch TEXT NOT NULL,
          action TEXT NOT NULL, reason TEXT, username TEXT, created_at TEXT);
        """)
        for branch_name in DEFAULT_BRANCHES:
            c.execute("INSERT OR IGNORE INTO branches_v13(name,active,created_by,created_at) VALUES(?,?,?,?)",
                      (branch_name,1,"system",datetime.now().isoformat(timespec='seconds')))
        # V13.5: planned delivery date follows the Order No. through the workflow.
        cols=[r[1] for r in c.execute("PRAGMA table_info(order_cycles)").fetchall()]
        if "delivery_date" not in cols:
            c.execute("ALTER TABLE order_cycles ADD COLUMN delivery_date TEXT")
        # V13.5.4: remember whether an allocation was explicitly saved.
        # This is required because a valid saved allocation can itself be 0.
        item_cols=[r[1] for r in c.execute("PRAGMA table_info(branch_order_items)").fetchall()]
        if "allocation_saved" not in item_cols:
            c.execute("ALTER TABLE branch_order_items ADD COLUMN allocation_saved INTEGER DEFAULT 0")
        # V13.6.1: preserve the exact generated DR PDF and receiving proof with the delivery record.
        delivery_cols=[r[1] for r in c.execute("PRAGMA table_info(deliveries_v2)").fetchall()]
        if "dr_pdf" not in delivery_cols:
            c.execute("ALTER TABLE deliveries_v2 ADD COLUMN dr_pdf BLOB")
        if "dr_filename" not in delivery_cols:
            c.execute("ALTER TABLE deliveries_v2 ADD COLUMN dr_filename TEXT")
        if "proof_data" not in delivery_cols:
            c.execute("ALTER TABLE deliveries_v2 ADD COLUMN proof_data BLOB")
        if "proof_mime" not in delivery_cols:
            c.execute("ALTER TABLE deliveries_v2 ADD COLUMN proof_mime TEXT")
        # Local test improvements: completion, allocation confirmation, and emergency archive.
        cycle_cols=[r[1] for r in c.execute("PRAGMA table_info(order_cycles)").fetchall()]
        if "completed_at" not in cycle_cols:
            c.execute("ALTER TABLE order_cycles ADD COLUMN completed_at TEXT")
        if "completed_by" not in cycle_cols:
            c.execute("ALTER TABLE order_cycles ADD COLUMN completed_by TEXT")
        cb_cols=[r[1] for r in c.execute("PRAGMA table_info(order_cycle_branches)").fetchall()]
        if "allocation_confirmation" not in cb_cols:
            c.execute("ALTER TABLE order_cycle_branches ADD COLUMN allocation_confirmation TEXT DEFAULT 'Pending'")
        if "allocation_confirmed_by" not in cb_cols:
            c.execute("ALTER TABLE order_cycle_branches ADD COLUMN allocation_confirmed_by TEXT")
        if "allocation_confirmed_at" not in cb_cols:
            c.execute("ALTER TABLE order_cycle_branches ADD COLUMN allocation_confirmed_at TEXT")
        if "allocation_concern" not in cb_cols:
            c.execute("ALTER TABLE order_cycle_branches ADD COLUMN allocation_concern TEXT")
        if "allocation_released" not in cb_cols:
            c.execute("ALTER TABLE order_cycle_branches ADD COLUMN allocation_released INTEGER DEFAULT 0")
        if "allocation_released_by" not in cb_cols:
            c.execute("ALTER TABLE order_cycle_branches ADD COLUMN allocation_released_by TEXT")
        if "allocation_released_at" not in cb_cols:
            c.execute("ALTER TABLE order_cycle_branches ADD COLUMN allocation_released_at TEXT")
        if "done_at" not in delivery_cols:
            c.execute("ALTER TABLE deliveries_v2 ADD COLUMN done_at TEXT")
        if "done_by" not in delivery_cols:
            c.execute("ALTER TABLE deliveries_v2 ADD COLUMN done_by TEXT")
        # Safe additive fields for Super Admin completion without branch receiving.
        if "done_without_receiving" not in delivery_cols:
            c.execute("ALTER TABLE deliveries_v2 ADD COLUMN done_without_receiving INTEGER DEFAULT 0")
        if "done_override_reason" not in delivery_cols:
            c.execute("ALTER TABLE deliveries_v2 ADD COLUMN done_override_reason TEXT")
        emer_cols=[r[1] for r in c.execute("PRAGMA table_info(emergency_orders)").fetchall()]
        if "archived" not in emer_cols:
            c.execute("ALTER TABLE emergency_orders ADD COLUMN archived INTEGER DEFAULT 0")
        if "archived_by" not in emer_cols:
            c.execute("ALTER TABLE emergency_orders ADD COLUMN archived_by TEXT")
        if "archived_at" not in emer_cols:
            c.execute("ALTER TABLE emergency_orders ADD COLUMN archived_at TEXT")
        if "archive_reason" not in emer_cols:
            c.execute("ALTER TABLE emergency_orders ADD COLUMN archive_reason TEXT")

def audit(action, details=""):
    with db_conn() as c:
        c.execute("INSERT INTO audit_v13(username,action,details,created_at) VALUES(?,?,?,?)",
                  (current_user,action,details,datetime.now().isoformat(timespec='seconds')))

def category_items(cat):
    path = TEMPLATE_FILES.get(cat)
    if not path or not path.exists(): return []
    wb=load_workbook(path,data_only=True); ws=wb.active; vals=[]
    if cat == "Pastries":
        for cell in ws[8][2:]:
            if cell.value not in (None, ""): vals.append(str(cell.value).strip())
    else:
        for r in range(9,ws.max_row+1):
            v=ws.cell(r,2).value
            if v not in (None,"","ITEM") and str(v).strip().upper()!="ITEM": vals.append(str(v).strip())
    vals = list(dict.fromkeys(vals))
    # Silvanas is an official Pastries ordering item. Keep it available even if an older
    # deployed Pastries template has not yet been replaced.
    if cat == "Pastries" and not any(v.strip().lower() == "silvanas" for v in vals):
        vals.append("Silvanas")
    return vals

def manila_now():
    return datetime.now(ZoneInfo("Asia/Manila"))

def order_lock_time(created_at):
    try:
        created = datetime.fromisoformat(str(created_at))
        d = created.date()
    except Exception:
        d = manila_now().date()
    return datetime(d.year, d.month, d.day, 21, 0, tzinfo=ZoneInfo("Asia/Manila"))

def is_order_locked(created_at):
    return manila_now() >= order_lock_time(created_at)

def branch_for_user():
    if st.session_state.get("branch"):
        return st.session_state.get("branch")
    rec = supabase_user(current_user) or load_users().get(current_user,{})
    return rec.get("branch") or current_user.replace('_',' ').title()

def available_branches(include_inactive=False):
    branches=[]
    try:
        with db_conn() as c:
            q="SELECT name FROM branches_v13" + ("" if include_inactive else " WHERE active=1") + " ORDER BY name"
            branches=[r[0] for r in c.execute(q).fetchall()]
    except Exception:
        branches=list(DEFAULT_BRANCHES)
    # Preserve any branch already assigned to an account so historical/user data is never orphaned.
    try:
        for rec in load_users().values():
            b=(rec.get('branch') or '').strip()
            if b and b != "Main Store / All Branches" and b not in branches:
                branches.append(b)
    except Exception:
        pass
    return sorted(dict.fromkeys(branches))

def branch_management():
    st.markdown("## Settings — Branch Management")
    st.caption("Add branches here once. Active branches automatically appear in User Management and Create Order Request.")
    with st.expander("➕ Add New Branch", expanded=True):
        name=st.text_input("Branch Name *", placeholder="e.g. Bayombong").strip()
        if st.button("Add Branch", type="primary", use_container_width=True):
            if not name:
                st.error("Branch name is required.")
            elif name.lower() == "main store / all branches":
                st.error("That name is reserved for Admin/Super Admin assignment.")
            else:
                try:
                    with db_conn() as c:
                        c.execute("INSERT INTO branches_v13(name,active,created_by,created_at) VALUES(?,?,?,?)",
                                  (name,1,current_user,datetime.now().isoformat(timespec='seconds')))
                    audit("BRANCH_ADDED",name)
                    st.success(f"{name} added successfully.")
                    st.rerun()
                except sqlite3.IntegrityError:
                    st.error("That branch already exists.")
    with db_conn() as c:
        rows=pd.read_sql_query("SELECT id,name,active,created_at FROM branches_v13 ORDER BY name",c)
    if rows.empty:
        st.info("No branches yet.")
        return
    show=rows.copy(); show['Status']=show['active'].map({1:'Active',0:'Inactive'})
    st.dataframe(show[['name','Status','created_at']].rename(columns={'name':'Branch','created_at':'Created At'}),use_container_width=True,hide_index=True)
    choice=st.selectbox("Select Branch to Manage", rows['name'].tolist())
    rec=rows[rows['name']==choice].iloc[0]
    if int(rec.active)==1:
        if st.button("Deactivate Selected Branch"):
            with db_conn() as c: c.execute("UPDATE branches_v13 SET active=0 WHERE id=?",(int(rec.id),))
            audit("BRANCH_DEACTIVATED",choice); st.success("Branch deactivated. Historical transactions are preserved."); st.rerun()
    else:
        if st.button("Reactivate Selected Branch"):
            with db_conn() as c: c.execute("UPDATE branches_v13 SET active=1 WHERE id=?",(int(rec.id),))
            audit("BRANCH_REACTIVATED",choice); st.success("Branch reactivated."); st.rerun()
    st.info("Branches are deactivated instead of deleted so old orders, DRs, receiving records, and history remain intact.")

def super_orders():
    st.markdown("## Orders")
    st.caption("Create an Order No., choose a category and the branches required to submit. The request will appear automatically on those branch users' dashboards.")
    emergency_order_admin()
    with st.expander("➕ Create New Order Request", expanded=True):
        c1,c2=st.columns(2)
        with c1: ono=st.text_input("Order No. *", placeholder="e.g. AB-ORD-2026-001").strip()
        with c2: cat=st.selectbox("Category *",ORDER_CATEGORIES,key="cyclecat")
        branches=available_branches()
        selected=st.multiselect("Branches required to order *",branches)
        d1,d2=st.columns(2)
        with d1:
            deadline=st.text_input("Submission deadline (optional)",placeholder="e.g. Sept 18, 2026 - 10:00 AM")
        with d2:
            planned_delivery_date=st.date_input("Delivery Date *", value=date.today())
        notes=st.text_area("Instructions / Notes (optional)")
        if st.button("Send Order Request",type="primary",use_container_width=True):
            if not ono: st.error("Order No. is required.")
            elif not selected: st.error("Select at least one branch.")
            else:
                try:
                    with db_conn() as c:
                        cur=c.execute("INSERT INTO order_cycles(order_no,category,created_by,created_at,deadline,notes,delivery_date) VALUES(?,?,?,?,?,?,?)",
                                      (ono,cat,current_user,datetime.now().isoformat(timespec='seconds'),deadline,notes,planned_delivery_date.isoformat()))
                        cid=cur.lastrowid
                        for b in selected: c.execute("INSERT INTO order_cycle_branches(cycle_id,branch) VALUES(?,?)",(cid,b))
                    audit("ORDER_REQUEST_SENT",f"{ono}; {cat}; {', '.join(selected)}")
                    st.success(f"Order request {ono} sent to {len(selected)} branch(es).")
                    st.rerun()
                except sqlite3.IntegrityError: st.error("That Order No. already exists. Please use a unique Order No.")
    st.markdown("### Order Requests & Submission Status")
    with db_conn() as c:
        cycles=pd.read_sql_query("SELECT id,order_no,category,status,created_at,deadline,delivery_date,completed_at FROM order_cycles WHERE COALESCE(status,'Open')!='Successful' ORDER BY id DESC",c)
    if cycles.empty: st.info("No order requests yet."); return
    labels=[f"{r.order_no} — {r.category}" for _,r in cycles.iterrows()]
    chosen=st.selectbox("Select Order Request",labels)
    cyc=cycles.iloc[labels.index(chosen)]
    if getattr(cyc, 'delivery_date', None):
        st.info(f"Planned Delivery Date: {cyc.delivery_date}")
    with db_conn() as c:
        status=pd.read_sql_query("SELECT id,branch,status,submitted_by,submitted_at FROM order_cycle_branches WHERE cycle_id=? ORDER BY branch",c,params=(int(cyc.id),))
    st.dataframe(status[['branch','status','submitted_by','submitted_at']].rename(columns={'branch':'Branch','status':'Status','submitted_by':'Submitted By','submitted_at':'Submitted At'}),use_container_width=True,hide_index=True)

    with st.expander("⚙️ Manage Branches on this Order Request", expanded=False):
        current_branches = status['branch'].astype(str).tolist()
        add_options = [b for b in available_branches() if b not in current_branches]
        add_branch = st.selectbox("Add Branch", [""] + add_options, key=f"add_branch_{int(cyc.id)}")
        if st.button("ADD BRANCH TO ORDER REQUEST", key=f"add_branch_btn_{int(cyc.id)}", use_container_width=True):
            if not add_branch:
                st.error("Select a branch to add.")
            else:
                with db_conn() as c:
                    c.execute("INSERT OR IGNORE INTO order_cycle_branches(cycle_id,branch,status) VALUES(?,?, 'Pending')", (int(cyc.id), add_branch))
                    c.execute("INSERT INTO order_branch_changes(cycle_id,branch,action,reason,username,created_at) VALUES(?,?,?,?,?,?)",
                              (int(cyc.id),add_branch,'Added','',current_user,manila_now().isoformat(timespec='seconds')))
                audit("ORDER_BRANCH_ADDED", f"{cyc.order_no}; {add_branch}")
                st.success(f"{add_branch} added to {cyc.order_no}."); st.rerun()

        removable = status['branch'].astype(str).tolist()
        remove_branch = st.selectbox("Remove / Cancel Branch", [""] + removable, key=f"remove_branch_{int(cyc.id)}")
        remove_reason = st.text_input("Reason for removing branch *", key=f"remove_reason_{int(cyc.id)}")
        if st.button("REMOVE / CANCEL BRANCH", key=f"remove_branch_btn_{int(cyc.id)}", use_container_width=True):
            if not remove_branch:
                st.error("Select a branch.")
            elif not remove_reason.strip():
                st.error("Reason is required.")
            else:
                row = status[status['branch'].astype(str)==remove_branch].iloc[0]
                with db_conn() as c:
                    has_items = c.execute("SELECT COUNT(*) FROM branch_order_items WHERE cycle_branch_id=?",(int(row.id),)).fetchone()[0]
                    if has_items or str(row.status) not in ('Pending',''):
                        c.execute("UPDATE order_cycle_branches SET status='Cancelled' WHERE id=?",(int(row.id),))
                        action='Cancelled'
                    else:
                        c.execute("DELETE FROM order_cycle_branches WHERE id=?",(int(row.id),))
                        action='Removed'
                    c.execute("INSERT INTO order_branch_changes(cycle_id,branch,action,reason,username,created_at) VALUES(?,?,?,?,?,?)",
                              (int(cyc.id),remove_branch,action,remove_reason.strip(),current_user,manila_now().isoformat(timespec='seconds')))
                audit("ORDER_BRANCH_REMOVED", f"{cyc.order_no}; {remove_branch}; {remove_reason.strip()}")
                st.success(f"{remove_branch} {action.lower()} from active order request."); st.rerun()

    # Submission state is permanent once a branch submits. Delivery/receiving is tracked separately.
    submitted=status[(status['status']!='Cancelled') & status['submitted_at'].notna() & status['submitted_at'].astype(str).str.strip().ne('')]
    pending=status[(status['status']=='Pending')]
    a,b=st.columns(2); a.metric("Submitted",len(submitted)); b.metric("Pending",len(pending))
    if not pending.empty: st.warning("Waiting for: "+", ".join(pending.branch.tolist()))
    # Regular submitted orders are one source of the consolidated order. Approved
    # Emergency Additional Orders are another source and MUST still appear even when
    # that branch has not submitted its regular order yet.
    submitted_ids=submitted.id.astype(int).tolist()
    active=status[status['status']!='Cancelled']
    active_ids=active.id.astype(int).tolist()

    if submitted_ids:
        submitted_marks=','.join('?'*len(submitted_ids))
        with db_conn() as c:
            df=pd.read_sql_query(f"""SELECT bi.id item_id,cb.id cycle_branch_id,cb.branch,bi.item,bi.ordered,bi.allocated,bi.allocation_saved,bi.delivered,bi.received
              FROM order_cycle_branches cb JOIN branch_order_items bi ON bi.cycle_branch_id=cb.id
              WHERE cb.id IN ({submitted_marks}) ORDER BY bi.item,cb.branch""",c,params=submitted_ids)
    else:
        df=pd.DataFrame(columns=['item_id','cycle_branch_id','branch','item','ordered','allocated','allocation_saved','delivered','received'])

    # Read approved emergency orders directly from the selected Order No./cycle.
    # Do not depend on whether the branch is Pending or Submitted: an approved
    # emergency order must be included in consolidation for this order cycle.
    with db_conn() as c:
        emer=pd.read_sql_query("""SELECT e.cycle_branch_id,cb.branch,e.item,
                 SUM(COALESCE(e.approved_qty,0)) approved_emergency
          FROM emergency_orders e
          JOIN order_cycle_branches cb ON cb.id=e.cycle_branch_id
          WHERE cb.cycle_id=?
            AND e.status IN ('Approved','Partially Approved')
            AND COALESCE(e.archived,0)=0
          GROUP BY e.cycle_branch_id,cb.branch,e.item""",c,params=(int(cyc.id),))

    # Add an order row for an approved emergency item that has no regular submitted
    # item row yet. Its original ordered quantity is 0; the approved emergency amount
    # becomes its consolidated quantity. This also makes pending branches visible.
    if not emer.empty:
        existing_keys=set(zip(pd.to_numeric(df.get('cycle_branch_id',pd.Series(dtype=float)),errors='coerce').fillna(-1).astype(int),df.get('item',pd.Series(dtype=str)).astype(str).str.strip().str.casefold()))
        extra=[]
        for _,er in emer.iterrows():
            key=(int(er['cycle_branch_id']),str(er['item']).strip().casefold())
            if key not in existing_keys:
                extra.append({'item_id':None,'cycle_branch_id':int(er['cycle_branch_id']),'branch':er['branch'],'item':er['item'],'ordered':0.0,'allocated':0.0,'allocation_saved':0,'delivered':0.0,'received':0.0})
        if extra:
            df=pd.concat([df,pd.DataFrame(extra)],ignore_index=True)

        # Normalize item text only for matching; preserve the displayed item name.
        df['_item_key']=df['item'].astype(str).str.strip().str.casefold()
        emer['_item_key']=emer['item'].astype(str).str.strip().str.casefold()
        emer_merge=emer[['cycle_branch_id','_item_key','approved_emergency']]
        df=df.merge(emer_merge,on=['cycle_branch_id','_item_key'],how='left').drop(columns=['_item_key'])
    else:
        df['approved_emergency']=0.0

    if df.empty:
        st.info("Consolidated order will appear after a regular order is submitted or an Emergency Order is approved.")
        return
    df['approved_emergency']=pd.to_numeric(df['approved_emergency'],errors='coerce').fillna(0.0)
    df['consolidated_order']=pd.to_numeric(df['ordered'],errors='coerce').fillna(0.0)+df['approved_emergency']

    pivot=df.pivot_table(index='item',columns='branch',values='consolidated_order',aggfunc='sum',fill_value=0)
    pivot['TOTAL']=pivot.sum(axis=1)
    st.markdown("### Consolidated Branch Orders")
    st.caption("Approved Emergency Additional Orders are automatically included. Original submitted orders remain preserved in history.")
    st.dataframe(pivot,use_container_width=True)
    st.download_button("Download Consolidated CSV",pivot.to_csv().encode(),f"{cyc.order_no}_Consolidated.csv","text/csv")
    st.markdown("### Prepare Delivery / Allocation")
    st.caption("The original branch orders are preserved. EDIT ALLOCATION changes only the To Deliver quantities.")
    alloc=df.copy()
    # If no allocation has been explicitly saved yet, start from the submitted order.
    # Once saved, even an allocation of 0 must remain 0 when the page is reopened.
    alloc['to_deliver']=alloc.apply(
        lambda r: float(r['allocated'] or 0) if int(r.get('allocation_saved',0) or 0)==1 else float(r['consolidated_order'] or 0), axis=1)
    branches_order=list(dict.fromkeys(alloc['branch'].tolist()))
    # Build the allocation grid from the actual saved order-item records.
    # This keeps every visible cell directly tied to a database item ID.
    items_order=list(dict.fromkeys(alloc['item'].astype(str).tolist()))
    edit_key=f"allocation_edit_{int(cyc.id)}"
    editing=bool(st.session_state.get(edit_key,False))

    if str(cyc.category) == 'Pastries':
        # Pastries: branches are rows, items are columns, TOTAL is a bottom row.
        matrix=alloc.pivot_table(index='branch',columns='item',values='to_deliver',aggfunc='sum',fill_value=0).reindex(index=branches_order,columns=items_order,fill_value=0)
        matrix.index.name='BRANCH'
        display_matrix=matrix.reset_index()
        editable_cols=items_order
        disabled_cols=['BRANCH']
    else:
        # Cups/Powders/Torani/Store: items are rows, branches are columns, TOTAL is on the right.
        matrix=alloc.pivot_table(index='item',columns='branch',values='to_deliver',aggfunc='sum',fill_value=0).reindex(index=items_order,columns=branches_order,fill_value=0)
        matrix.index.name='ITEM'
        matrix['TOTAL']=matrix.sum(axis=1)
        display_matrix=matrix.reset_index()
        editable_cols=branches_order
        disabled_cols=['ITEM','TOTAL']

    def allocation_with_totals(frame):
        live=frame.copy()
        if str(cyc.category)=='Pastries':
            total={'BRANCH':'TOTAL'}
            for col in editable_cols:
                total[col]=float(pd.to_numeric(live[col],errors='coerce').fillna(0).sum())
            return pd.concat([live,pd.DataFrame([total])],ignore_index=True)
        live['TOTAL']=live[editable_cols].apply(pd.to_numeric,errors='coerce').fillna(0).sum(axis=1)
        return live

    if not editing:
        st.dataframe(allocation_with_totals(display_matrix),use_container_width=True,hide_index=True)
        if st.button("✏️ EDIT ALLOCATION",type="primary",use_container_width=True):
            st.session_state[edit_key]=True
            st.rerun()
        edited_matrix=display_matrix
        live=allocation_with_totals(edited_matrix)
    else:
        st.info("Edit the To Deliver quantities below. Original branch orders remain locked and unchanged.")
        # Keep the editor inside a form so typing a quantity does not rerun the page
        # and rebuild emergency-only cells from their original emergency quantity.
        # The edited values are submitted together only when SAVE ALLOCATION is clicked.
        with st.form(f"allocation_form_{int(cyc.id)}", clear_on_submit=False):
            edited_matrix=st.data_editor(display_matrix,hide_index=True,use_container_width=True,
                disabled=disabled_cols,
                column_config={col:st.column_config.NumberColumn(col,min_value=0.0,step=1.0) for col in editable_cols},
                key=f"allocation_grid_{int(cyc.id)}")
            live=allocation_with_totals(edited_matrix)
            st.markdown("#### Updated Allocation Totals")
            st.dataframe(live,use_container_width=True,hide_index=True)
            csave,ccancel=st.columns(2)
            save_clicked=csave.form_submit_button("SAVE ALLOCATION",type="primary",use_container_width=True)
            cancel_clicked=ccancel.form_submit_button("CANCEL EDIT",use_container_width=True)
        if cancel_clicked:
            st.session_state[edit_key]=False
            st.rerun()
        if save_clicked:
            changes=[]; missing=[]

            # V13.5.6: save by POSITION against parallel ID/ordered matrices.
            # The editor's row/column positions are created from these same matrices,
            # so saving never tries to match user-visible branch/item text.
            if str(cyc.category)=='Pastries':
                id_matrix=alloc.pivot_table(index='branch',columns='item',values='item_id',aggfunc='first').reindex(index=branches_order,columns=items_order)
                ordered_matrix=alloc.pivot_table(index='branch',columns='item',values='ordered',aggfunc='sum',fill_value=0).reindex(index=branches_order,columns=items_order,fill_value=0)
                for row_pos in range(len(branches_order)):
                    for col_pos in range(len(items_order)):
                        qty=float(pd.to_numeric(pd.Series([edited_matrix.iloc[row_pos, col_pos+1]]),errors='coerce').fillna(0).iloc[0])
                        raw_id=id_matrix.iloc[row_pos,col_pos]
                        if pd.isna(raw_id):
                            missing.append(f"row {row_pos+1}, column {col_pos+1}")
                            continue
                        item_id=int(raw_id)
                        changes.append((qty,item_id,items_order[col_pos]))
            else:
                id_matrix=alloc.pivot_table(index='item',columns='branch',values='item_id',aggfunc='first').reindex(index=items_order,columns=branches_order)
                ordered_matrix=alloc.pivot_table(index='item',columns='branch',values='ordered',aggfunc='sum',fill_value=0).reindex(index=items_order,columns=branches_order,fill_value=0)
                for row_pos in range(len(items_order)):
                    for col_pos in range(len(branches_order)):
                        qty=float(pd.to_numeric(pd.Series([edited_matrix.iloc[row_pos, col_pos+1]]),errors='coerce').fillna(0).iloc[0])
                        raw_id=id_matrix.iloc[row_pos,col_pos]
                        if pd.isna(raw_id):
                            missing.append(f"row {row_pos+1}, column {col_pos+1}")
                            continue
                        item_id=int(raw_id)
                        changes.append((qty,item_id,items_order[row_pos]))
            # Save all existing cells. Missing matrix cells are legitimate when a branch originally
            # submitted no row for an item. Create that item only when Admin allocates > 0; otherwise
            # a missing zero cell needs no database row. This avoids position/matching failures.
            with db_conn() as c:
                for qty,item_id,item in changes:
                    c.execute("UPDATE branch_order_items SET allocated=?, allocation_saved=1 WHERE id=?",(qty,item_id))
                if missing:
                    if str(cyc.category)=='Pastries':
                        for row_pos in range(len(branches_order)):
                            for col_pos in range(len(items_order)):
                                raw_id=id_matrix.iloc[row_pos,col_pos]
                                if pd.isna(raw_id):
                                    qty=float(pd.to_numeric(pd.Series([edited_matrix.iloc[row_pos,col_pos+1]]),errors='coerce').fillna(0).iloc[0])
                                    if qty>0:
                                        cbid=int(alloc.loc[alloc['branch']==branches_order[row_pos],'cycle_branch_id'].iloc[0])
                                        c.execute("INSERT OR IGNORE INTO branch_order_items(cycle_branch_id,item,ordered,allocated,allocation_saved,delivered,received) VALUES(?,?,?,?,1,0,0)",(cbid,items_order[col_pos],0,qty))
                    else:
                        for row_pos in range(len(items_order)):
                            for col_pos in range(len(branches_order)):
                                raw_id=id_matrix.iloc[row_pos,col_pos]
                                if pd.isna(raw_id):
                                    qty=float(pd.to_numeric(pd.Series([edited_matrix.iloc[row_pos,col_pos+1]]),errors='coerce').fillna(0).iloc[0])
                                    if qty>0:
                                        cbid=int(alloc.loc[alloc['branch']==branches_order[col_pos],'cycle_branch_id'].iloc[0])
                                        c.execute("INSERT OR IGNORE INTO branch_order_items(cycle_branch_id,item,ordered,allocated,allocation_saved,delivered,received) VALUES(?,?,?,?,1,0,0)",(cbid,items_order[row_pos],0,qty))
            audit("ALLOCATION_SAVED",str(cyc.order_no))
            st.session_state[edit_key]=False
            st.success("Allocation saved. The saved quantities now reflect as the Allocation / To Deliver for this order.")
            st.rerun()

    # Excel allocation download follows the category orientation used on screen.
    export=allocation_with_totals(edited_matrix)
    out=io.BytesIO()
    with pd.ExcelWriter(out,engine='openpyxl') as writer:
        export.to_excel(writer,index=False,sheet_name='Allocation',startrow=4)
        ws=writer.book['Allocation']
        ws['A1']='APAYAO BREW — ALLOCATION'
        ws['A2']=f"Order No.: {cyc.order_no}"
        ws['A3']=f"Category: {cyc.category} | Delivery Date: {getattr(cyc,'delivery_date','') or ''}"
        for cell in ws[5]:
            cell.font=cell.font.copy(bold=True)
        ws.freeze_panes='B6'
        for col in ws.columns:
            width=min(max(len(str(x.value or '')) for x in col)+2,32)
            ws.column_dimensions[col[0].column_letter].width=max(width,12)
    st.download_button("⬇️ DOWNLOAD ALLOCATION EXCEL",out.getvalue(),f"{cyc.order_no}_{cyc.category}_Allocation.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)

    st.markdown("#### Upload Allocation Excel")
    st.caption("Upload the same Allocation Excel format. Only To Deliver / allocation quantities are updated; original orders remain unchanged.")
    allocation_upload=st.file_uploader("UPLOAD ALLOCATION EXCEL",type=["xlsx"],key=f"allocation_upload_{int(cyc.id)}")
    if allocation_upload is not None:
        try:
            up=pd.read_excel(allocation_upload,sheet_name="Allocation",skiprows=4)
            up=up.loc[:,~up.columns.astype(str).str.startswith("Unnamed")].copy()
            if "TOTAL" in up.columns: up=up.drop(columns=["TOTAL"])
            if str(cyc.category)=="Pastries":
                up=up[up["BRANCH"].astype(str).str.upper()!="TOTAL"].copy()
            st.dataframe(up,use_container_width=True,hide_index=True)
            if st.button("APPLY / SAVE UPLOADED ALLOCATION",type="primary",use_container_width=True,key=f"apply_alloc_upload_{int(cyc.id)}"):
                updates=[]; errors=[]
                with db_conn() as c:
                    cbmap={str(r[1]).strip():int(r[0]) for r in c.execute("SELECT id,branch FROM order_cycle_branches WHERE cycle_id=? AND status!='Cancelled'",(int(cyc.id),)).fetchall()}
                    if str(cyc.category)=="Pastries":
                        if "BRANCH" not in up.columns: errors.append("BRANCH column is missing.")
                        else:
                            for _,rr in up.iterrows():
                                br=str(rr["BRANCH"]).strip()
                                if br not in cbmap: errors.append(f"Unknown branch: {br}"); continue
                                for item in [x for x in up.columns if x!="BRANCH"]:
                                    q=pd.to_numeric(pd.Series([rr[item]]),errors="coerce").iloc[0]
                                    if pd.isna(q) or float(q)<0: errors.append(f"Invalid quantity: {br} / {item}"); continue
                                    updates.append((cbmap[br],str(item).strip(),float(q)))
                    else:
                        if "ITEM" not in up.columns: errors.append("ITEM column is missing.")
                        else:
                            for _,rr in up.iterrows():
                                item=str(rr["ITEM"]).strip()
                                for br in [x for x in up.columns if x!="ITEM"]:
                                    if str(br).strip() not in cbmap: errors.append(f"Unknown branch: {br}"); continue
                                    q=pd.to_numeric(pd.Series([rr[br]]),errors="coerce").iloc[0]
                                    if pd.isna(q) or float(q)<0: errors.append(f"Invalid quantity: {item} / {br}"); continue
                                    updates.append((cbmap[str(br).strip()],item,float(q)))
                    if errors:
                        st.error("Upload not applied. " + " | ".join(list(dict.fromkeys(errors))[:10]))
                    else:
                        for cbid,item,qty in updates:
                            row=c.execute("SELECT id FROM branch_order_items WHERE cycle_branch_id=? AND lower(trim(item))=lower(trim(?))",(cbid,item)).fetchone()
                            if row: c.execute("UPDATE branch_order_items SET allocated=?,allocation_saved=1 WHERE id=?",(qty,int(row[0])))
                            elif qty>0: c.execute("INSERT INTO branch_order_items(cycle_branch_id,item,ordered,allocated,allocation_saved,delivered,received) VALUES(?,?,0,?,1,0,0)",(cbid,item,qty))
                # The allocation transaction must COMMIT before audit opens another
                # SQLite connection. Auditing inside `with db_conn()` locked the DB.
                if not errors:
                    audit("ALLOCATION_EXCEL_UPLOADED",str(cyc.order_no))
                    st.success("Uploaded allocation saved. The uploaded quantities now reflect as the Allocation / To Deliver for this order.")
                    st.rerun()
        except Exception as exc:
            st.error(f"Could not process or save this allocation file: {exc}")

    with db_conn() as c:
        conf=pd.read_sql_query("SELECT branch,COALESCE(allocation_released,0) released,COALESCE(allocation_confirmation,'Pending') confirmation FROM order_cycle_branches WHERE cycle_id=? AND status!='Cancelled' ORDER BY branch",c,params=(int(cyc.id),))
    st.markdown("#### Branch Allocation Confirmation")
    conf_show=conf.copy(); conf_show['Release Status']=conf_show['released'].map(lambda x:'Released' if int(x or 0)==1 else 'Not Released')
    st.dataframe(conf_show[['branch','Release Status','confirmation']].rename(columns={'branch':'Branch','confirmation':'Confirmation'}),use_container_width=True,hide_index=True)
    if st.button("📤 RELEASE ALLOCATION FOR BRANCH CONFIRMATION",type="primary",use_container_width=True,key=f"release_alloc_{int(cyc.id)}"):
        with db_conn() as c:
            c.execute("UPDATE order_cycle_branches SET allocation_released=1,allocation_released_by=?,allocation_released_at=? WHERE cycle_id=? AND status!='Cancelled'",(current_user,manila_now().isoformat(timespec='seconds'),int(cyc.id)))
        audit("ALLOCATION_RELEASED",str(cyc.order_no)); st.success("Allocation released. Branches can now confirm it."); st.rerun()
    all_confirmed = (not conf.empty) and conf['released'].fillna(0).astype(int).eq(1).all() and conf['confirmation'].eq('Confirmed').all()
    if not all_confirmed:
        st.info("DR can proceed after all active branches confirm their allocation. Branches may also report an allocation concern.")
    if st.button("CONFIRM DELIVERY / SEND FOR RECEIVING",type="primary",use_container_width=True,disabled=not all_confirmed):
        with db_conn() as c:
            rows=pd.read_sql_query("SELECT bi.cycle_branch_id,bi.item,bi.ordered,bi.allocated,bi.allocation_saved FROM branch_order_items bi JOIN order_cycle_branches cb ON cb.id=bi.cycle_branch_id WHERE cb.cycle_id=? AND cb.allocation_confirmation='Confirmed'",c,params=(int(cyc.id),))
            for _,r in rows.iterrows():
                # Respect a saved allocation exactly, including a deliberate 0.
                qty=float(r['allocated'] or 0) if int(r['allocation_saved'] or 0)==1 else float(r['ordered'] or 0)
                c.execute("UPDATE branch_order_items SET delivered=? WHERE cycle_branch_id=? AND item=?",(qty,int(r.cycle_branch_id),r['item']))
            for cbid in rows.cycle_branch_id.unique():
                dno=f"DR-{cyc.order_no}-{int(cbid)}"
                c.execute("INSERT OR IGNORE INTO deliveries_v2(delivery_no,cycle_branch_id,created_by,created_at) VALUES(?,?,?,?)",(dno,int(cbid),current_user,datetime.now().isoformat(timespec='seconds')))
                # Keep the branch order submission status as Submitted.
                # The delivery row itself carries the separate For Receiving status.
                c.execute("UPDATE order_cycle_branches SET status='Submitted' WHERE id=?",(int(cbid),))
        audit("DELIVERY_PREPARED",str(cyc.order_no)); st.success("Delivery confirmed and sent to the selected branches for receiving.")

    st.divider()
    if st.button("✅ MARK ORDER REQUEST SUCCESSFUL",use_container_width=True,key=f"successful_{int(cyc.id)}"):
        with db_conn() as c:
            c.execute("UPDATE order_cycles SET status='Successful',completed_by=?,completed_at=? WHERE id=?",(current_user,manila_now().isoformat(timespec='seconds'),int(cyc.id)))
        audit("ORDER_REQUEST_SUCCESSFUL",str(cyc.order_no))
        st.success("Order Request moved out of the active list and kept in Transaction History.")
        st.rerun()

def staff_orders():
    st.markdown("## Order Requests")
    branch=branch_for_user(); st.caption(f"Branch: {branch}")
    with db_conn() as c:
        req=pd.read_sql_query("""SELECT cb.id cycle_branch_id,oc.order_no,oc.category,oc.deadline,oc.delivery_date,oc.notes,cb.status,oc.created_at
          FROM order_cycle_branches cb JOIN order_cycles oc ON oc.id=cb.cycle_id
          WHERE cb.branch=? AND COALESCE(oc.status,'Open')!='Successful' ORDER BY cb.id DESC""",c,params=(branch,))
    if req.empty: st.info("No order request has been sent to your branch yet."); return
    pending=req[req.status=='Pending']
    if not pending.empty: st.warning(f"You have {len(pending)} order request(s) waiting for submission.")
    labels=[f"{r.order_no} — {r.category} — {r.status}" for _,r in req.iterrows()]
    chosen=st.selectbox("Select Order Request",labels)
    rec=req.iloc[labels.index(chosen)]
    st.markdown(f"### {rec.order_no}")
    st.write(f"**Branch:** {branch}  |  **Category:** {rec.category}")
    if rec.deadline: st.write(f"**Deadline:** {rec.deadline}")
    if getattr(rec,'delivery_date',None): st.write(f"**Delivery Date:** {rec.delivery_date}")
    if rec.notes: st.info(str(rec.notes))
    locked=is_order_locked(rec.created_at)
    st.caption(f"Normal ordering locks at 9:00 PM Philippine time. Cutoff: {order_lock_time(rec.created_at).strftime('%b %d, %Y %I:%M %p')}")

    # Show submitted/no-order/cancelled records first.
    if rec.status!='Pending':
        with db_conn() as c:
            old=pd.read_sql_query("SELECT item,ordered,allocated,delivered,received FROM branch_order_items WHERE cycle_branch_id=?",c,params=(int(rec.cycle_branch_id),))
        if str(rec.status)=='No Order': st.success("NO ORDER was confirmed for this request.")
        elif str(rec.status)=='Cancelled': st.warning("This branch was removed/cancelled from this order request by Super Admin.")
        if not old.empty:
            old['Remaining']=old['ordered']-old['delivered']; st.dataframe(old,use_container_width=True,hide_index=True)

    # Branch must confirm the saved allocation before Super Admin can proceed to DR.
    with db_conn() as c:
        alloc_confirm=pd.read_sql_query("SELECT item,ordered,allocated,allocation_saved FROM branch_order_items WHERE cycle_branch_id=? ORDER BY id",c,params=(int(rec.cycle_branch_id),))
        confrow=c.execute("SELECT COALESCE(allocation_confirmation,'Pending'),COALESCE(allocation_concern,''),COALESCE(allocation_released,0) FROM order_cycle_branches WHERE id=?",(int(rec.cycle_branch_id),)).fetchone()
    released=int(confrow[2] or 0)==1 if confrow else False
    has_saved=(not alloc_confirm.empty) and alloc_confirm['allocation_saved'].fillna(0).astype(int).eq(1).any()
    if released:
        st.markdown("### Allocation for Confirmation")
        if alloc_confirm.empty:
            st.info("NO ORDER was submitted. Allocated Qty: 0")
            av=pd.DataFrame(columns=['item','ordered','allocated','Remaining'])
        else:
            av=alloc_confirm.copy(); av['Remaining']=pd.to_numeric(av['ordered'],errors='coerce').fillna(0)-pd.to_numeric(av['allocated'],errors='coerce').fillna(0)
        st.dataframe(av[['item','ordered','allocated','Remaining']].rename(columns={'item':'Item','ordered':'Ordered Qty','allocated':'Allocated Qty'}),use_container_width=True,hide_index=True)
        state=confrow[0] if confrow else 'Pending'
        st.write(f"**Allocation Status:** {state}")
        if state!='Confirmed':
            ca,cb=st.columns(2)
            if ca.button("CONFIRM ALLOCATION",type="primary",use_container_width=True,key=f"confirm_alloc_{int(rec.cycle_branch_id)}"):
                with db_conn() as c: c.execute("UPDATE order_cycle_branches SET allocation_confirmation='Confirmed',allocation_confirmed_by=?,allocation_confirmed_at=?,allocation_concern=NULL WHERE id=?",(current_user,manila_now().isoformat(timespec='seconds'),int(rec.cycle_branch_id)))
                audit("ALLOCATION_CONFIRMED",f"{rec.order_no}; {branch}"); st.success("Allocation confirmed."); st.rerun()
            concern=st.text_input("Allocation concern / correction needed",key=f"alloc_concern_{int(rec.cycle_branch_id)}")
            if cb.button("REPORT ALLOCATION CONCERN",use_container_width=True,key=f"concern_alloc_{int(rec.cycle_branch_id)}"):
                if not concern.strip(): st.error("Enter the concern first.")
                else:
                    with db_conn() as c: c.execute("UPDATE order_cycle_branches SET allocation_confirmation='Concern',allocation_concern=?,allocation_confirmed_by=?,allocation_confirmed_at=? WHERE id=?",(concern.strip(),current_user,manila_now().isoformat(timespec='seconds'),int(rec.cycle_branch_id)))
                    audit("ALLOCATION_CONCERN",f"{rec.order_no}; {branch}; {concern.strip()}"); st.success("Concern sent to Super Admin."); st.rerun()

    # Emergency additions are available after the 9 PM cutoff for active requests, including already-submitted orders.
    if locked and str(rec.status)!='Cancelled':
        st.warning("Normal ordering is LOCKED because the 9:00 PM Philippine-time cutoff has passed.")
        items=category_items(rec.category)
        if items:
            with st.expander("🚨 EMERGENCY ADDITIONAL ORDER", expanded=(str(rec.status)=='Pending')):
                st.caption("This does not change your original order. Enter only the extra quantities needed.")
                emer=pd.DataFrame({'Item':items,'ADDITIONAL QTY':[0.0]*len(items)})
                emer_edit=st.data_editor(emer,hide_index=True,use_container_width=True,disabled=['Item'],
                    column_config={'ADDITIONAL QTY':st.column_config.NumberColumn(min_value=0.0,step=1.0)},
                    key=f"emer_editor_{int(rec.cycle_branch_id)}")
                reason=st.text_area("Reason for Emergency Additional Order *",key=f"emer_reason_{int(rec.cycle_branch_id)}")
                if st.button("SUBMIT EMERGENCY ADDITIONAL ORDER",type="primary",use_container_width=True,key=f"emer_submit_{int(rec.cycle_branch_id)}"):
                    active=emer_edit[pd.to_numeric(emer_edit['ADDITIONAL QTY'],errors='coerce').fillna(0)>0]
                    if active.empty: st.error("Enter at least one additional quantity greater than 0.")
                    elif not reason.strip(): st.error("Reason is required for an emergency additional order.")
                    else:
                        now=manila_now().isoformat(timespec='seconds')
                        with db_conn() as c:
                            for _,r in active.iterrows():
                                c.execute("INSERT INTO emergency_orders(cycle_branch_id,item,requested_qty,status,reason,requested_by,requested_at) VALUES(?,?,?,?,?,?,?)",
                                          (int(rec.cycle_branch_id),str(r['Item']),float(r['ADDITIONAL QTY']),'Pending Approval',reason.strip(),current_user,now))
                        audit("EMERGENCY_ORDER_SUBMITTED",f"{rec.order_no}; {branch}; {reason.strip()}")
                        st.success("Emergency Additional Order submitted for Super Admin approval."); st.rerun()
        with db_conn() as c:
            er=pd.read_sql_query("SELECT item,requested_qty,approved_qty,status,reason,requested_at FROM emergency_orders WHERE cycle_branch_id=? AND COALESCE(archived,0)=0 ORDER BY id DESC",c,params=(int(rec.cycle_branch_id),))
        if not er.empty:
            st.markdown("#### Emergency Order History"); st.dataframe(er,use_container_width=True,hide_index=True)
        return

    # Before the 9 PM cutoff, a branch may revise its own submitted order.
    # Order No., category and branch stay protected; only item quantities can change.
    if str(rec.status) in ('Submitted', 'No Order'):
        edit_key=f"edit_order_{int(rec.cycle_branch_id)}"
        cedit,cnoedit=st.columns(2)
        with cedit:
            if st.button("EDIT ORDER",type="primary",use_container_width=True,key=f"edit_btn_{int(rec.cycle_branch_id)}"):
                st.session_state[edit_key]=True
                st.rerun()
        with cnoedit:
            if str(rec.status)=='Submitted' and st.button("CHANGE TO NO ORDER",use_container_width=True,key=f"edit_no_btn_{int(rec.cycle_branch_id)}"):
                st.session_state[f"edit_confirm_no_{int(rec.cycle_branch_id)}"]=True

        edit_no_key=f"edit_confirm_no_{int(rec.cycle_branch_id)}"
        if st.session_state.get(edit_no_key):
            st.warning("Change this submitted order to NO ORDER? The previous quantities will remain in the audit history.")
            yy,nn=st.columns(2)
            if yy.button("YES, CHANGE TO NO ORDER",type="primary",use_container_width=True,key=f"edit_yes_no_{int(rec.cycle_branch_id)}"):
                now=manila_now().isoformat(timespec='seconds')
                with db_conn() as c:
                    before=pd.read_sql_query("SELECT item,ordered FROM branch_order_items WHERE cycle_branch_id=? ORDER BY id",c,params=(int(rec.cycle_branch_id),))
                    c.execute("UPDATE branch_order_items SET ordered=0 WHERE cycle_branch_id=?",(int(rec.cycle_branch_id),))
                    c.execute("UPDATE order_cycle_branches SET status='No Order',submitted_by=?,submitted_at=? WHERE id=?",(current_user,now,int(rec.cycle_branch_id)))
                audit("BRANCH_ORDER_EDITED_TO_NO_ORDER",json.dumps({"order_no":str(rec.order_no),"branch":branch,"previous":before.to_dict('records')},default=str))
                st.session_state.pop(edit_no_key,None); st.session_state.pop(edit_key,None)
                st.success("Order changed to NO ORDER."); st.rerun()
            if nn.button("CANCEL",use_container_width=True,key=f"edit_cancel_no_{int(rec.cycle_branch_id)}"):
                st.session_state.pop(edit_no_key,None); st.rerun()

        if st.session_state.get(edit_key):
            items=category_items(rec.category)
            if not items:
                st.error("No item masterlist found for this category."); return
            with db_conn() as c:
                existing=pd.read_sql_query("SELECT item,ordered FROM branch_order_items WHERE cycle_branch_id=?",c,params=(int(rec.cycle_branch_id),))
            existing_map={str(r.item):float(r.ordered or 0) for _,r in existing.iterrows()}
            edit_df=pd.DataFrame({'Item':items,'MY ORDER':[existing_map.get(i,0.0) for i in items]})
            revised=st.data_editor(edit_df,hide_index=True,use_container_width=True,disabled=['Item'],
                column_config={'MY ORDER':st.column_config.NumberColumn(min_value=0.0,step=1.0)},
                key=f"edit_order_editor_{int(rec.cycle_branch_id)}")
            st.caption("You may revise quantities until 9:00 PM Philippine time. Enter 0 when you intentionally need none of an item.")
            save_col,cancel_col=st.columns(2)
            with save_col:
                save_edit=st.button("SAVE & RESUBMIT ORDER",type="primary",use_container_width=True,key=f"save_edit_{int(rec.cycle_branch_id)}")
            with cancel_col:
                cancel_edit=st.button("CANCEL EDIT",use_container_width=True,key=f"cancel_edit_{int(rec.cycle_branch_id)}")
            if cancel_edit:
                st.session_state.pop(edit_key,None); st.rerun()
            if save_edit:
                if is_order_locked(rec.created_at):
                    st.error("The 9:00 PM cutoff has passed. Use Emergency Additional Order instead.")
                    st.session_state.pop(edit_key,None); st.rerun()
                vals=pd.to_numeric(revised['MY ORDER'],errors='coerce')
                if vals.isna().any():
                    st.error("Order incomplete. Please enter a quantity for every item.")
                else:
                    now=manila_now().isoformat(timespec='seconds')
                    before_records=existing.to_dict('records')
                    after_records=[]
                    with db_conn() as c:
                        for idx,r in revised.iterrows():
                            qty=float(vals.iloc[idx])
                            after_records.append({'item':str(r['Item']),'ordered':qty})
                            c.execute("""INSERT INTO branch_order_items(cycle_branch_id,item,ordered) VALUES(?,?,?)
                                       ON CONFLICT(cycle_branch_id,item) DO UPDATE SET ordered=excluded.ordered""",
                                      (int(rec.cycle_branch_id),str(r['Item']),qty))
                        c.execute("UPDATE order_cycle_branches SET status='Submitted',submitted_by=?,submitted_at=? WHERE id=?",(current_user,now,int(rec.cycle_branch_id)))
                    audit("BRANCH_ORDER_EDITED",json.dumps({"order_no":str(rec.order_no),"branch":branch,"before":before_records,"after":after_records},default=str))
                    st.session_state.pop(edit_key,None)
                    st.success("Order updated and resubmitted successfully."); st.rerun()
        return

    if rec.status!='Pending':
        return

    items=category_items(rec.category)
    if not items: st.error("No item masterlist found for this category."); return
    base=pd.DataFrame({'Item':items,'MY ORDER':[None]*len(items)})
    edited=st.data_editor(base,hide_index=True,use_container_width=True,disabled=['Item'],
        column_config={'MY ORDER':st.column_config.NumberColumn(min_value=0.0,step=1.0,help='Enter 0 if you have no order for this item.')},
        key=f"normal_order_{int(rec.cycle_branch_id)}")
    st.caption("Every item is required for a normal order. Or use NO ORDER if the branch needs nothing from this request.")

    csubmit,cno=st.columns(2)
    with csubmit:
        review_clicked=st.button("REVIEW & SUBMIT",type="primary",use_container_width=True,key=f"review_{int(rec.cycle_branch_id)}")
    with cno:
        no_clicked=st.button("NO ORDER",use_container_width=True,key=f"no_order_{int(rec.cycle_branch_id)}")
    if review_clicked:
        missing=edited['MY ORDER'].isna()
        if missing.any():
            st.error("Order incomplete. Please enter a quantity for all items, or use NO ORDER.")
            st.warning("Missing: "+", ".join(edited.loc[missing,'Item'].astype(str).tolist()))
        else:
            st.markdown("### Review Order"); st.dataframe(edited,use_container_width=True,hide_index=True)
            st.session_state[f"review_order_{int(rec.cycle_branch_id)}"] = edited.to_dict('records')
    if no_clicked:
        st.session_state[f"confirm_no_order_{int(rec.cycle_branch_id)}"]=True

    nokey=f"confirm_no_order_{int(rec.cycle_branch_id)}"
    if st.session_state.get(nokey):
        st.warning("Are you sure your branch has NO ORDER for this order request?")
        y,n=st.columns(2)
        if y.button("YES, CONFIRM NO ORDER",type="primary",use_container_width=True,key=f"yes_no_{int(rec.cycle_branch_id)}"):
            now=manila_now().isoformat(timespec='seconds')
            with db_conn() as c:
                c.execute("UPDATE order_cycle_branches SET status='No Order',submitted_by=?,submitted_at=? WHERE id=?",(current_user,now,int(rec.cycle_branch_id)))
            st.session_state.pop(nokey,None); audit("BRANCH_NO_ORDER",f"{rec.order_no}; {branch}")
            st.success("No Order confirmed. This counts as a completed response."); st.rerun()
        if n.button("CANCEL",use_container_width=True,key=f"cancel_no_{int(rec.cycle_branch_id)}"):
            st.session_state.pop(nokey,None); st.rerun()

    key=f"review_order_{int(rec.cycle_branch_id)}"
    if key in st.session_state:
        if st.button("CONFIRM & SUBMIT ORDER",type="primary",use_container_width=True,key=f"confirm_submit_{int(rec.cycle_branch_id)}"):
            rows=pd.DataFrame(st.session_state[key]); now=manila_now().isoformat(timespec='seconds')
            with db_conn() as c:
                for _,r in rows.iterrows(): c.execute("INSERT INTO branch_order_items(cycle_branch_id,item,ordered) VALUES(?,?,?)",(int(rec.cycle_branch_id),r['Item'],float(r['MY ORDER'])))
                c.execute("UPDATE order_cycle_branches SET status='Submitted',submitted_by=?,submitted_at=? WHERE id=?",(current_user,now,int(rec.cycle_branch_id)))
            st.session_state.pop(key,None); audit("BRANCH_ORDER_SUBMITTED",f"{rec.order_no}; {branch}")
            st.success("Order submitted successfully. Original quantities are preserved."); st.rerun()

def my_orders():
    staff_orders()

def receive_delivery():
    st.markdown("## Receive Delivery")
    branch=branch_for_user(); st.caption(f"Branch: {branch}")
    with db_conn() as c:
        dels=pd.read_sql_query("""SELECT d.id,d.delivery_no,d.cycle_branch_id,oc.order_no,oc.category,d.created_at,d.dr_pdf,d.dr_filename
          FROM deliveries_v2 d JOIN order_cycle_branches cb ON cb.id=d.cycle_branch_id
          JOIN order_cycles oc ON oc.id=cb.cycle_id WHERE cb.branch=? AND d.status='For Receiving' ORDER BY d.id DESC""",c,params=(branch,))
    if dels.empty: st.info("No delivery waiting for receiving."); return
    label=st.selectbox("Select Delivery",dels.delivery_no.tolist()); rec=dels[dels.delivery_no==label].iloc[0]
    st.markdown(f"### {rec.order_no} • {label}")

    # Exact DR generated and saved by Super Admin.
    if rec.dr_pdf is not None:
        pdf_bytes=bytes(rec.dr_pdf)
        st.download_button("📄 VIEW / DOWNLOAD DELIVERY RECEIPT",data=pdf_bytes,
            file_name=(rec.dr_filename or f"DR_{safe_filename(label)}_{safe_filename(branch)}.pdf"),
            mime="application/pdf",use_container_width=True)
    else:
        st.info("The DR PDF has not been saved to this delivery yet. Super Admin can regenerate it and click SAVE DR DETAILS.")

    with db_conn() as c:
        items=pd.read_sql_query("SELECT item,ordered,delivered,received FROM branch_order_items WHERE cycle_branch_id=? ORDER BY id",c,params=(int(rec.cycle_branch_id),))
    for col in ['ordered','delivered']:
        items[col]=pd.to_numeric(items[col],errors='coerce').fillna(0)
    show=items[['item','ordered','delivered']].copy(); show.columns=['Item','Ordered','DR Qty / To Deliver']
    st.markdown("#### Finalized Delivery Reference")
    st.dataframe(show,use_container_width=True,hide_index=True)
    st.markdown("#### Actual Receiving / Tally")
    st.caption("Enter the quantity physically received for every item. Enter 0 if none was received. DR Qty is read-only.")

    rows=[]; incomplete=False
    h1,h2,h3,h4=st.columns([3,1,1,1])
    h1.markdown("**Item**"); h2.markdown("**DR Qty**"); h3.markdown("**Actual Received**"); h4.markdown("**Variance**")
    for idx,r in items.iterrows():
        c1,c2,c3,c4=st.columns([3,1,1,1])
        c1.write(str(r['item'])); c2.write(f"{float(r['delivered']):g}")
        val=c3.number_input("Actual Received",min_value=0.0,step=1.0,value=None,key=f"recv_qty_{int(rec.id)}_{idx}",label_visibility="collapsed")
        if val is None:
            incomplete=True; variance=None; c4.write("—")
        else:
            variance=float(val)-float(r['delivered']); c4.write(f"{variance:+g}" if variance else "0")
        rows.append({'Item':r['item'],'DR / Delivered':float(r['delivered']),'Actual Received':val,'Variance':variance})

    entry=pd.DataFrame(rows)
    hasvar=entry['Variance'].notna().any() and (entry['Variance'].dropna()!=0).any()
    remarks=st.text_area("Explanation / reason *" if hasvar else "Remarks (optional)",key=f"recv_remarks_{int(rec.id)}")
    proof=st.file_uploader("Proof/photo *",type=['png','jpg','jpeg','pdf'],key=f"recv_proof_{int(rec.id)}") if hasvar else None
    if hasvar: st.warning("Variance detected. Explanation and proof are required before submission.")

    if st.button("REVIEW RECEIVING",type="primary",use_container_width=True):
        if incomplete:
            st.error("Receiving incomplete. Enter Actual Received for every item. Enter 0 if none was received."); return
        if hasvar and (not remarks.strip() or proof is None):
            st.error("Variance detected. Explanation and proof are required."); return
        st.session_state[f"recv_{int(rec.id)}"]={
            'rows':entry.to_dict('records'),'remarks':remarks,
            'proof_name':proof.name if proof else '',
            'proof_data':proof.getvalue() if proof else None,
            'proof_mime':proof.type if proof else ''}
        st.success("Review complete. Confirm below to submit the receiving record.")

    key=f"recv_{int(rec.id)}"
    if key in st.session_state:
        review=pd.DataFrame(st.session_state[key]['rows'])
        st.markdown("##### Review")
        st.dataframe(review,use_container_width=True,hide_index=True)
        if st.button("CONFIRM & SUBMIT RECEIVING",type="primary",use_container_width=True):
            data=st.session_state[key]; rows=pd.DataFrame(data['rows'])
            hv=(pd.to_numeric(rows['Variance'],errors='coerce').fillna(0)!=0).any()
            with db_conn() as c:
                for _,r in rows.iterrows():
                    c.execute("UPDATE branch_order_items SET received=? WHERE cycle_branch_id=? AND item=?",(float(r['Actual Received']),int(rec.cycle_branch_id),r['Item']))
                status='Received - With Variance' if hv else 'Received - No Variance'
                c.execute("UPDATE deliveries_v2 SET status=?,received_by=?,received_at=?,remarks=?,proof_name=?,proof_data=?,proof_mime=? WHERE id=?",
                          (status,current_user,datetime.now().isoformat(timespec='seconds'),data['remarks'],data['proof_name'],data['proof_data'],data['proof_mime'],int(rec.id)))
                c.execute("UPDATE order_cycle_branches SET status=? WHERE id=?",(status,int(rec.cycle_branch_id)))
            st.session_state.pop(key,None); audit("DELIVERY_RECEIVED",f"{label}; variance={hv}"); st.success("Receiving submitted. It is now waiting for Super Admin review."); st.rerun()

def receiving_variances_admin():
    st.markdown("## Receiving & Variances")
    st.caption("Select an Order No. to complete multiple branches together. Received and not-received branches have separate actions.")
    with db_conn() as c:
        summary=pd.read_sql_query("""SELECT d.id delivery_id,d.cycle_branch_id,d.delivery_no,oc.order_no,oc.category,
          oc.delivery_date,cb.branch,d.status,d.received_by,d.received_at,d.remarks,d.proof_name,
          COALESCE(SUM(bi.delivered),0) allocated,COALESCE(SUM(bi.received),0) actual_received
          FROM deliveries_v2 d JOIN order_cycle_branches cb ON cb.id=d.cycle_branch_id
          JOIN order_cycles oc ON oc.id=cb.cycle_id
          LEFT JOIN branch_order_items bi ON bi.cycle_branch_id=cb.id
          WHERE COALESCE(d.status,'')!='Done Delivery'
          GROUP BY d.id,d.cycle_branch_id,d.delivery_no,oc.order_no,oc.category,oc.delivery_date,
            cb.branch,d.status,d.received_by,d.received_at,d.remarks,d.proof_name
          ORDER BY d.id DESC""",c)
    if summary.empty:
        st.info("No active receiving or variance transactions.")
        return
    order_no=st.selectbox("Select Order No.",summary.order_no.dropna().unique().tolist(),key="bulk_done_order")
    view=summary[summary.order_no==order_no].copy()
    view['Receiving Status']=view.status.map(lambda x: 'Submitted' if str(x).startswith('Received') else 'Not Submitted')
    view['Actual Received']=view.apply(lambda r: float(r.actual_received or 0) if str(r.status).startswith('Received') else None,axis=1)
    view['Variance']=view.apply(lambda r: float(r.actual_received or 0)-float(r.allocated or 0) if str(r.status).startswith('Received') else None,axis=1)
    st.dataframe(view[['branch','delivery_no','allocated','Actual Received','Variance','Receiving Status','status']].rename(columns={
        'branch':'Branch','delivery_no':'DR No.','allocated':'DR Qty','status':'Delivery Status'}),use_container_width=True,hide_index=True)
    received=view[view.status.astype(str).str.startswith('Received')]
    waiting=view[~view.status.astype(str).str.startswith('Received')]
    st.markdown("### Complete branches with submitted receiving")
    if received.empty:
        st.info("No branches have submitted receiving for this order.")
    else:
        received_options={f"{r.branch} — {r.delivery_no} (#{int(r.delivery_id)})":int(r.delivery_id) for _,r in received.iterrows()}
        all_received=st.checkbox("Select All Received",key=f"all_received_{order_no}")
        selected_received=list(received_options) if all_received else st.multiselect("Select received branches",list(received_options),key=f"bulk_received_{order_no}")
        confirm_received=st.checkbox("Confirm completion of the selected received deliveries",key=f"confirm_received_{order_no}")
        if st.button("✅ CONFIRM SELECTED DONE DELIVERY",type="primary",disabled=not selected_received or not confirm_received,key=f"complete_received_{order_no}"):
            ids=[received_options[label] for label in selected_received]
            completed=[]
            with db_conn() as c:
                for did in ids:
                    cur=c.execute("""UPDATE deliveries_v2 SET status='Done Delivery',done_by=?,done_at=?,
                      done_without_receiving=0,done_override_reason=NULL
                      WHERE id=? AND status LIKE 'Received%'""",
                      (current_user,manila_now().isoformat(timespec='seconds'),did))
                    if cur.rowcount:
                        cbid=c.execute("SELECT cycle_branch_id FROM deliveries_v2 WHERE id=?",(did,)).fetchone()[0]
                        c.execute("UPDATE order_cycle_branches SET status='Done Delivery' WHERE id=?",(cbid,))
                        completed.append(did)
            for did in completed: audit("DONE_DELIVERY",f"Order={order_no}; delivery_id={did}; bulk")
            st.success(f"{len(completed)} received delivery/deliveries filed under Deliveries.")
            st.rerun()
    st.markdown("### Complete branches without receiving")
    if waiting.empty:
        st.info("No branches are waiting to submit receiving for this order.")
    elif current_role!='Super Admin':
        st.info("Only Super Admin can complete a delivery without submitted receiving.")
    else:
        waiting_options={f"{r.branch} — {r.delivery_no} (#{int(r.delivery_id)})":int(r.delivery_id) for _,r in waiting.iterrows()}
        all_waiting=st.checkbox("Select All Not Received",key=f"all_waiting_{order_no}")
        selected_waiting=list(waiting_options) if all_waiting else st.multiselect("Select branches without receiving",list(waiting_options),key=f"bulk_waiting_{order_no}")
        reason=st.text_area("Required reason for completion without receiving *",key=f"bulk_reason_{order_no}")
        confirm_waiting=st.checkbox("I confirm these branches have not submitted receiving",key=f"confirm_waiting_{order_no}")
        if st.button("⚠️ MARK SELECTED DONE WITHOUT RECEIVING",disabled=not selected_waiting,key=f"complete_waiting_{order_no}"):
            if not reason.strip(): st.error("A reason is required.")
            elif not confirm_waiting: st.error("Please confirm the override.")
            else:
                ids=[waiting_options[label] for label in selected_waiting]
                completed=[]
                with db_conn() as c:
                    for did in ids:
                        cur=c.execute("""UPDATE deliveries_v2 SET status='Done Delivery',done_by=?,done_at=?,
                          done_without_receiving=1,done_override_reason=?
                          WHERE id=? AND status NOT LIKE 'Received%' AND status!='Done Delivery'""",
                          (current_user,manila_now().isoformat(timespec='seconds'),reason.strip(),did))
                        if cur.rowcount:
                            cbid=c.execute("SELECT cycle_branch_id FROM deliveries_v2 WHERE id=?",(did,)).fetchone()[0]
                            c.execute("UPDATE order_cycle_branches SET status='Done Delivery' WHERE id=?",(cbid,))
                            completed.append(did)
                for did in completed: audit("DONE_DELIVERY_WITHOUT_RECEIVING",f"Order={order_no}; delivery_id={did}; bulk; reason={reason.strip()}")
                st.success(f"{len(completed)} delivery/deliveries completed without receiving and filed under Deliveries.")
                st.rerun()
    st.divider()
    st.markdown("### View active delivery items")
    detail_options={f"{r.branch} — {r.delivery_no} (#{int(r.delivery_id)})":int(r.delivery_id) for _,r in view.iterrows()}
    choice=st.selectbox("View branch delivery",list(detail_options),key=f"active_detail_{order_no}")
    rec=view[view.delivery_id==detail_options[choice]].iloc[0]
    with db_conn() as c:
        detail=pd.read_sql_query("SELECT item,ordered,delivered,received FROM branch_order_items WHERE cycle_branch_id=? ORDER BY id",c,params=(int(rec.cycle_branch_id),))
    detail=detail.rename(columns={'item':'Item','ordered':'Ordered Qty','delivered':'DR Qty','received':'Actual Received'})
    if not str(rec.status).startswith('Received'):
        detail['Actual Received']='Not Submitted'
        detail['Variance']='Not Submitted'
    else:
        detail['Variance']=pd.to_numeric(detail['Actual Received'],errors='coerce')-pd.to_numeric(detail['DR Qty'],errors='coerce')
    st.dataframe(detail,use_container_width=True,hide_index=True)
    if rec.remarks: st.write(f"**Remarks:** {rec.remarks}")
    if rec.proof_name: st.write(f"**Proof:** {rec.proof_name}")


def deliveries_history():
    st.markdown("## Deliveries")
    st.caption("Completed deliveries remain available with individual delivered items and receiving details.")
    with db_conn() as c:
        df=pd.read_sql_query("""SELECT d.id delivery_id,d.cycle_branch_id,oc.created_at order_date,
          oc.delivery_date,d.done_at completed_at,oc.order_no,oc.category,cb.branch,d.delivery_no,
          d.status,d.received_by,d.received_at,d.done_by,d.remarks,d.proof_name,
          d.dr_pdf,d.dr_filename,d.proof_data,d.proof_mime,
          COALESCE(d.done_without_receiving,0) done_without_receiving,d.done_override_reason,
          COALESCE(SUM(bi.ordered),0) ordered,COALESCE(SUM(bi.delivered),0) delivered,
          COALESCE(SUM(bi.received),0) received
          FROM deliveries_v2 d JOIN order_cycle_branches cb ON cb.id=d.cycle_branch_id
          JOIN order_cycles oc ON oc.id=cb.cycle_id
          LEFT JOIN branch_order_items bi ON bi.cycle_branch_id=cb.id
          WHERE d.status='Done Delivery'
          GROUP BY d.id,d.cycle_branch_id,oc.created_at,oc.delivery_date,d.done_at,oc.order_no,
            oc.category,cb.branch,d.delivery_no,d.status,d.received_by,d.received_at,d.done_by,
            d.remarks,d.proof_name,d.dr_pdf,d.dr_filename,d.proof_data,d.proof_mime,
            d.done_without_receiving,d.done_override_reason
          ORDER BY COALESCE(d.done_at,d.received_at,d.created_at) DESC""",c)
    if df.empty:
        st.info("No completed deliveries yet.")
        return
    df['Receiving Status']=df.apply(lambda r: 'Not Submitted — Super Admin Override' if int(r.done_without_receiving or 0) else 'Received',axis=1)
    df['Actual Received Display']=df.apply(lambda r: None if int(r.done_without_receiving or 0) else float(r.received or 0),axis=1)
    df['Variance']=df.apply(lambda r: None if int(r.done_without_receiving or 0) else float(r.received or 0)-float(r.delivered or 0),axis=1)
    f1,f2,f3=st.columns(3)
    cat=f1.selectbox("Category",['All']+sorted(df.category.dropna().astype(str).unique()),key="del_cat")
    br=f2.selectbox("Branch",['All']+sorted(df.branch.dropna().astype(str).unique()),key="del_branch")
    search=f3.text_input("Order No.",key="del_order_search")
    view=df.copy()
    if cat!='All': view=view[view.category==cat]
    if br!='All': view=view[view.branch==br]
    if search.strip(): view=view[view.order_no.astype(str).str.contains(search.strip(),case=False,na=False)]
    table=view[['order_date','delivery_date','order_no','category','branch','delivery_no','ordered','delivered','Actual Received Display','Variance','Receiving Status','done_by','done_override_reason','completed_at']].copy()
    table.columns=['Order Date','Delivery Date','Order No.','Category','Branch','DR No.','Ordered Qty','Delivered Qty','Actual Received','Variance','Receiving Status','Done By','Override Reason','Done Date']
    st.dataframe(table,use_container_width=True,hide_index=True)
    out=io.BytesIO()
    with pd.ExcelWriter(out,engine='openpyxl') as writer: table.to_excel(writer,index=False,sheet_name='Deliveries')
    st.download_button("⬇️ DOWNLOAD FILTERED DELIVERIES EXCEL",out.getvalue(),"Apayao_Brew_Deliveries.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True)
    if view.empty: return
    st.markdown("### VIEW DELIVERY DETAILS")
    options={f"{r.order_no} — {r.branch} — {r.delivery_no} (#{int(r.delivery_id)})":int(r.delivery_id) for _,r in view.iterrows()}
    pick=st.selectbox("Select completed delivery",list(options),key="completed_delivery_detail")
    rec=view[view.delivery_id==options[pick]].iloc[0]
    st.write(f"**Order:** {rec.order_no} | **Category:** {rec.category} | **Branch:** {rec.branch} | **DR:** {rec.delivery_no}")
    st.write(f"**Completed by:** {rec.done_by or 'Not recorded'} | **Completed at:** {rec.completed_at or 'Not recorded'}")
    if int(rec.done_without_receiving or 0):
        st.warning("Completed by Super Admin without branch receiving. Actual Received and Variance were not submitted.")
        st.write(f"**Override reason:** {rec.done_override_reason or 'Not recorded'}")
    else:
        st.write(f"**Received by:** {rec.received_by or 'Not recorded'} | **Received at:** {rec.received_at or 'Not recorded'}")
    with db_conn() as c:
        items=pd.read_sql_query("SELECT item,ordered,delivered,received FROM branch_order_items WHERE cycle_branch_id=? ORDER BY id",c,params=(int(rec.cycle_branch_id),))
    items=items.rename(columns={'item':'Item','ordered':'Ordered Qty','delivered':'Delivered / DR Qty','received':'Actual Received'})
    if int(rec.done_without_receiving or 0):
        items['Actual Received']='Not Submitted'
        items['Variance']='Not Submitted'
    else:
        items['Variance']=pd.to_numeric(items['Actual Received'],errors='coerce')-pd.to_numeric(items['Delivered / DR Qty'],errors='coerce')
    st.dataframe(items[['Item','Ordered Qty','Delivered / DR Qty','Actual Received','Variance']],use_container_width=True,hide_index=True)
    if rec.remarks: st.write(f"**Receiving remarks:** {rec.remarks}")
    if rec.proof_name: st.write(f"**Proof:** {rec.proof_name}")
    if rec.proof_data is not None:
        st.download_button("⬇️ DOWNLOAD RECEIVING PROOF",bytes(rec.proof_data),file_name=rec.proof_name or 'receiving_proof',mime=rec.proof_mime or 'application/octet-stream',key=f"proof_{int(rec.delivery_id)}")
    if rec.dr_pdf is not None:
        st.download_button("📄 DOWNLOAD DELIVERY RECEIPT",bytes(rec.dr_pdf),file_name=rec.dr_filename or f"DR_{int(rec.delivery_id)}.pdf",mime='application/pdf',key=f"dr_history_{int(rec.delivery_id)}")


def database_backup():
    st.markdown("## Database Backup")
    st.caption("Download a copy of the current live Apayao Brew SQLite database. Keep this file private because it contains operational records.")
    if current_role != "Super Admin":
        st.error("Database backup is available to Super Admin only.")
        return
    if not DB_FILE.exists():
        st.error("Database file was not found.")
        return
    try:
        backup_bytes = DB_FILE.read_bytes()
        stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        st.success(f"Backup ready • {len(backup_bytes) / 1024:.1f} KB")
        st.download_button(
            "⬇️ DOWNLOAD LIVE DATABASE BACKUP",
            data=backup_bytes,
            file_name=f"Apayao_Brew_Database_Backup_{stamp}.db",
            mime="application/octet-stream",
            use_container_width=True,
            type="primary",
        )
        st.info("This downloads a copy only. It does not delete or change your live database.")
    except Exception as exc:
        st.error(f"Could not prepare the database backup: {exc}")


def transaction_history():
    st.markdown("## Transaction History")
    branch=None if current_role in ('Admin','Super Admin') else branch_for_user()
    with db_conn() as c:
        sql="""SELECT oc.order_no,oc.category,cb.branch,cb.status,cb.submitted_by,cb.submitted_at,d.delivery_no,d.status delivery_status,d.received_by,d.received_at
          FROM order_cycle_branches cb JOIN order_cycles oc ON oc.id=cb.cycle_id LEFT JOIN deliveries_v2 d ON d.cycle_branch_id=cb.id"""
        if branch: df=pd.read_sql_query(sql+" WHERE cb.branch=? ORDER BY cb.id DESC",c,params=(branch,))
        else: df=pd.read_sql_query(sql+" ORDER BY cb.id DESC",c)
    if df.empty:
        st.info("No transactions yet.")
    else:
        st.dataframe(df,use_container_width=True,hide_index=True)

init_v13_db()

def render_manual_dr_generator():
    brand_logo, brand_text, brand_user = st.columns([0.55, 5.2, 1.6])
    with brand_logo:
        if BUILTIN_LOGO.exists():
            st.image(str(BUILTIN_LOGO), width=64)
    with brand_text:
        st.markdown(
            """
            <div style="padding-top:5px">
                <div class="ab-brandname">APAYAO BREW</div>
                <div class="ab-brandtag">THE PERFECT BLEND.</div>
            </div>
            """,
            unsafe_allow_html=True
        )
    with brand_user:
        st.markdown(
            f"""<div class="ab-user"><b>{current_user}</b><br>{current_role}</div>""",
            unsafe_allow_html=True
        )

    st.markdown(
        """
        <div class="ab-hero">
            <div class="ab-welcome">Welcome!</div>
            <h1>Delivery Receipt Generator</h1>
            <p>Create, manage, and download your delivery receipts with ease.</p>
            <div class="ab-perfect">THE PERFECT BLEND.</div>
        </div>
        """,
        unsafe_allow_html=True
    )

    st.markdown('<div class="ab-sectiontitle">Select Category</div>', unsafe_allow_html=True)
    st.markdown('<div class="ab-sectionhint">Choose a category to get started.</div>', unsafe_allow_html=True)

    category = st.radio(
        "Category",
        CATEGORIES,
        horizontal=True,
        label_visibility="collapsed"
    )

    if category == "Pastries":
        st.info("Pastries template: branches are listed down the left side and products are across the top.")
    else:
        st.info(f"{category} template: branches are across the top row and items are listed down the first column.")

    st.markdown(
        """
        <div class="ab-sectiontitle">Generate Delivery Receipt</div>
        <div class="ab-sectionhint">Follow the steps below to create your delivery receipt.</div>
        <div class="ab-steps">
            <div class="ab-step"><div class="ab-stepnum">1</div><b>Download Template</b><span>Get the Excel template for your selected category.</span></div>
            <div class="ab-step"><div class="ab-stepnum">2</div><b>Upload Filled Excel</b><span>Upload the completed template with branch quantities.</span></div>
            <div class="ab-step"><div class="ab-stepnum">3</div><b>Generate PDF</b><span>Complete the required DR information and generate the receipts.</span></div>
            <div class="ab-step"><div class="ab-stepnum">4</div><b>Download</b><span>Download one combined PDF or separate branch files.</span></div>
        </div>
        """,
        unsafe_allow_html=True
    )

    template_path = CATEGORY_TEMPLATES.get(category)
    if template_path and template_path.exists():
        st.download_button(
            f"⬇️ Download {category} Excel Template",
            data=template_path.read_bytes(),
            file_name=template_path.name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True
        )
    elif category == "Store":
        st.info("Store template has not been added yet.")

    st.markdown("#### Upload Completed Template")
    uploaded = st.file_uploader("Upload Completed Excel Template", type=["xlsx","xls"], label_visibility="collapsed")

    if uploaded:
        try:
            raw = pd.read_excel(uploaded, header=None)
            branches = detect_matrix_by_category(raw, category)

            found_dr = txt(find_label_value(raw, ["DR NO", "DR NO."]))
            found_origin = txt(find_label_value(raw, ["FROM", "STORE ORIGIN"])) or "MAIN STORE"
            found_delivered = txt(find_label_value(raw, ["DELIVERED BY", "DELIVERY BY"]))
            found_date = excel_date(find_label_value(raw, ["DATE", "DELIVERY DATE"])) or date.today()

            st.success(f"Excel read successfully. Found {len(branches)} branch(es).")

            st.subheader("3. Delivery Information")
            c1, c2 = st.columns(2)

            with c1:
                base_dr = st.text_input("DR Number *", value=found_dr, placeholder="Example: 265")
                store_origin = st.text_input("Store Origin *", value=found_origin)
                prepared_by = st.text_input("Prepared By *", value="")
            with c2:
                delivered_by = st.text_input("Delivered By *", value=found_delivered)
                delivery_date = st.date_input("Delivery Date *", value=found_date)
                approved_by = st.text_input("Approved By *", value="Glady Clemente/ Coleen Navaro")

            preview = []
            for b in branches:
                preview.append({
                    "Branch": b["branch"],
                    "Category": category,
                    "DR": dr_no_text(base_dr, delivery_date, b["branch"]),
                    "Items with Quantity": sum(1 for _, q in b["items"] if q != 0),
                })

            st.subheader("Preview")
            st.dataframe(pd.DataFrame(preview), use_container_width=True, hide_index=True)

            logo_bytes = None

            required_fields = {
                "DR Number": base_dr.strip(),
                "Store Origin": store_origin.strip(),
                "Prepared By": prepared_by.strip(),
                "Delivered By": delivered_by.strip(),
                "Approved By": approved_by.strip(),
            }
            missing_fields = [name for name, value in required_fields.items() if not value]

            if missing_fields:
                st.warning(
                    "Complete all required DR details first: "
                    + ", ".join(missing_fields)
                )

            if st.button(
                "Generate Delivery Receipts",
                type="primary",
                use_container_width=True,
                disabled=bool(missing_fields)
            ):
                combined = build_combined_pdf(
                    branches, category, base_dr, store_origin,
                    prepared_by, delivered_by, delivery_date,
                    approved_by, current_user, logo_bytes
                )

                log_activity(
                    current_user,
                    "DR_GENERATED",
                    f"Category={category}; DR={base_dr}; Branches={len(branches)}"
                )
                st.success(f"Generated {len(branches)} branch receipt(s).")

                st.download_button(
                    "Download ONE Combined PDF",
                    data=combined,
                    file_name=f"Apayao_Brew_{safe_filename(category)}_Delivery_Receipts.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )

                # Also keep separate PDFs as an optional ZIP.
                zip_buf = io.BytesIO()
                with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
                    for b in branches:
                        pdf = build_single_pdf(
                            b, category, base_dr, store_origin,
                            prepared_by, delivered_by, delivery_date,
                            approved_by, current_user, logo_bytes
                        )
                        zf.writestr(
                            f"DR_{safe_filename(category)}_{safe_filename(b['branch'])}.pdf",
                            pdf
                        )
                zip_buf.seek(0)

                st.download_button(
                    "Optional: Download Separate PDFs (ZIP)",
                    data=zip_buf.getvalue(),
                    file_name=f"Apayao_Brew_{safe_filename(category)}_Separate_DRs.zip",
                    mime="application/zip",
                    use_container_width=True
                )

        except Exception as e:
            st.error(f"Could not read this Excel file: {e}")
    else:
        st.info("Select the category and upload your Excel file.")

    st.markdown(
        """
        <div class="ab-footer">
            <div><b>APAYAO BREW</b> &nbsp; | &nbsp; THE PERFECT BLEND.</div>
            <div>Delivery Receipt Generator</div>
        </div>
        """,
        unsafe_allow_html=True
    )


def render_dr_generator():
    st.markdown("## Delivery Receipt Generator")
    st.caption("Generate DRs from confirmed branch allocations. Select one, several, or all branches and preview before downloading.")
    tab_auto, tab_manual = st.tabs(["AUTOMATIC DR", "MANUAL EXCEL DR"])
    with tab_auto:
        with db_conn() as c:
            cycles=pd.read_sql_query("""SELECT oc.id,oc.order_no,oc.category,oc.delivery_date,oc.created_at
              FROM order_cycles oc WHERE EXISTS (SELECT 1 FROM order_cycle_branches cb JOIN deliveries_v2 d ON d.cycle_branch_id=cb.id WHERE cb.cycle_id=oc.id) ORDER BY oc.id DESC""",c)
        if cycles.empty:
            st.info("No confirmed deliveries yet. Complete branch allocation confirmation first.")
        else:
            labels=[f"{r.order_no} — {r.category}" for _,r in cycles.iterrows()]
            chosen=st.selectbox("Order No. *",labels,key="auto_dr_order")
            cyc=cycles.iloc[labels.index(chosen)]
            with db_conn() as c:
                branches=pd.read_sql_query("""SELECT cb.id cycle_branch_id,cb.branch,d.id delivery_id,d.delivery_no,d.status
                  FROM order_cycle_branches cb JOIN deliveries_v2 d ON d.cycle_branch_id=cb.id
                  WHERE cb.cycle_id=? ORDER BY cb.branch""",c,params=(int(cyc.id),))
            names=branches['branch'].astype(str).tolist()
            select_all=st.checkbox("SELECT ALL BRANCHES",value=False,key=f"dr_all_{int(cyc.id)}")
            selected_names=names if select_all else st.multiselect("Branches *",names,key=f"dr_branches_{int(cyc.id)}")
            default_date=date.today()
            try:
                if str(cyc.delivery_date or '').strip(): default_date=pd.to_datetime(cyc.delivery_date).date()
            except Exception: pass
            c1,c2=st.columns(2)
            with c1:
                base_dr=st.text_input("DR Base No. *",value=str(cyc.order_no),key=f"multi_drno_{int(cyc.id)}")
                store_origin=st.text_input("Store Origin *",value="MAIN STORE",key=f"multi_origin_{int(cyc.id)}")
                prepared_by=st.text_input("Prepared By *",value=current_user,key=f"multi_prep_{int(cyc.id)}")
                delivery_date=st.date_input("Delivery Date *",value=default_date,key=f"multi_date_{int(cyc.id)}")
            with c2:
                delivered_by=st.text_input("Delivered By / Driver *",key=f"multi_driver_{int(cyc.id)}")
                vehicle=st.text_input("Vehicle / Plate No.",key=f"multi_vehicle_{int(cyc.id)}")
                approved_by=st.text_input("Approved By *",value="Glady Clemente/ Coleen Navaro",key=f"multi_approve_{int(cyc.id)}")
                remarks=st.text_area("Remarks",key=f"multi_remarks_{int(cyc.id)}",height=70)
            required={'Branch':selected_names,'DR Base No.':base_dr.strip(),'Store Origin':store_origin.strip(),'Prepared By':prepared_by.strip(),'Delivered By / Driver':delivered_by.strip(),'Approved By':approved_by.strip()}
            missing=[k for k,v in required.items() if not v]
            if missing: st.warning("Complete required DR details: "+", ".join(missing))
            preview_key=f"dr_preview_{int(cyc.id)}"
            if st.button("👁️ PREVIEW SELECTED DRs",type="primary",use_container_width=True,disabled=bool(missing),key=f"preview_btn_{int(cyc.id)}"):
                st.session_state[preview_key]=True
            if st.session_state.get(preview_key) and not missing:
                payload=[]
                st.markdown("### DR Preview")
                for bname in selected_names:
                    br=branches[branches.branch==bname].iloc[0]
                    with db_conn() as c:
                        items=pd.read_sql_query("SELECT item,ordered,delivered FROM branch_order_items WHERE cycle_branch_id=? ORDER BY id",c,params=(int(br.cycle_branch_id),))
                    items['ordered']=pd.to_numeric(items['ordered'],errors='coerce').fillna(0); items['delivered']=pd.to_numeric(items['delivered'],errors='coerce').fillna(0)
                    payload.append({'branch':bname,'items':[(r['item'],float(r['delivered'])) for _,r in items.iterrows()]})
                    with st.expander(f"{bname} — {dr_no_text(base_dr,delivery_date,bname)}",expanded=(len(selected_names)<=3)):
                        st.dataframe(items.rename(columns={'item':'Item','ordered':'Ordered Qty','delivered':'To Deliver'}),use_container_width=True,hide_index=True)
                pdf=build_combined_pdf(payload,cyc.category,base_dr,store_origin,prepared_by,delivered_by,delivery_date,approved_by,current_user,None)
                st.info(f"Preview ready for {len(selected_names)} branch(es). The combined PDF prints one branch per page.")
                if vehicle: st.write(f"**Vehicle / Plate No.:** {vehicle}")
                if remarks: st.write(f"**Remarks:** {remarks}")
                if st.button("SAVE DR DETAILS FOR SELECTED BRANCHES",use_container_width=True,key=f"save_multi_{int(cyc.id)}"):
                    try:
                        with db_conn() as c:
                            for bname in selected_names:
                                br=branches[branches.branch==bname].iloc[0]
                                final_no=dr_no_text(base_dr,delivery_date,bname)
                                single_payload=next(x for x in payload if x['branch']==bname)
                                single_pdf=build_single_pdf(single_payload,cyc.category,base_dr,store_origin,prepared_by,delivered_by,delivery_date,approved_by,current_user,None)
                                c.execute("UPDATE deliveries_v2 SET delivery_no=?,remarks=?,dr_pdf=?,dr_filename=? WHERE id=?",(final_no,remarks.strip(),sqlite3.Binary(single_pdf),f"DR_{safe_filename(final_no)}_{safe_filename(bname)}.pdf",int(br.delivery_id)))
                        audit("DR_DETAILS_SAVED",f"Order={cyc.order_no}; Branches={', '.join(selected_names)}; Vehicle={vehicle}")
                        st.success("DR details saved for all selected branches.")
                    except sqlite3.IntegrityError:
                        st.error("A generated DR No. is already in use. Change the DR Base No. and preview again.")
                st.download_button("⬇️ DOWNLOAD ALL SELECTED DRs — ONE PDF",data=pdf,file_name=f"DR_{safe_filename(base_dr)}_ALL_SELECTED.pdf",mime="application/pdf",use_container_width=True)
    with tab_manual:
        st.info("Use this only for deliveries that did not originate from an Order Request.")
        render_manual_dr_generator()



# ========================= STOCK COUNT REQUESTS (SUPABASE) =========================
def _supabase_config():
    try:
        url = str(st.secrets.get("SUPABASE_URL", "")).rstrip("/")
        key = str(st.secrets.get("SUPABASE_SECRET_KEY", ""))
    except Exception:
        url, key = "", ""
    return url, key


def _sb_request(table, method="GET", params=None, payload=None):
    url, key = _supabase_config()
    if not url or not key:
        raise RuntimeError("Supabase is not configured. Add SUPABASE_URL and SUPABASE_SECRET_KEY to Streamlit Secrets.")
    endpoint = f"{url}/rest/v1/{table}"
    if params:
        endpoint += "?" + urllib.parse.urlencode(params, doseq=True)
    headers = {
        "apikey": key,
        "Authorization": f"Bearer {key}",
        "Content-Type": "application/json",
        "Accept": "application/json",
        "Prefer": "return=representation",
    }
    data = None if payload is None else json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(endpoint, data=data, headers=headers, method=method)
    try:
        with urllib.request.urlopen(req, timeout=30) as resp:
            body = resp.read().decode("utf-8")
            return json.loads(body) if body else []
    except urllib.error.HTTPError as e:
        detail = e.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"Supabase error {e.code}: {detail}")


def _sb_get(table, **filters):
    params = {"select": "*"}
    for k, v in filters.items():
        params[k] = f"eq.{v}"
    return _sb_request(table, "GET", params=params)


def _sb_insert(table, rows):
    return _sb_request(table, "POST", payload=rows)


def _sb_patch(table, filters, values):
    params = {k: f"eq.{v}" for k, v in filters.items()}
    return _sb_request(table, "PATCH", params=params, payload=values)


def emergency_order_admin():
    st.markdown("### 🚨 Emergency Additional Orders")
    with db_conn() as c:
        q=pd.read_sql_query("""SELECT e.id,e.item,e.requested_qty,e.approved_qty,e.status,e.reason,e.requested_by,e.requested_at,e.reviewed_by,e.reviewed_at,e.review_note,COALESCE(e.archived,0) archived,
          cb.branch,oc.order_no,oc.category FROM emergency_orders e
          JOIN order_cycle_branches cb ON cb.id=e.cycle_branch_id JOIN order_cycles oc ON oc.id=cb.cycle_id
          ORDER BY e.id DESC""",c)
    if q.empty:
        st.info("No emergency additional orders yet."); return
    pending=q[(q.status=='Pending Approval') & (q.archived==0)]
    if pending.empty:
        st.success("No pending emergency orders.")
    else:
        st.dataframe(pending[['order_no','branch','category','item','requested_qty','status','reason','requested_by','requested_at']],use_container_width=True,hide_index=True)
        labels=[f"#{int(r.id)} — {r.order_no} — {r.branch} — {r['item']} — Qty {r.requested_qty:g}" for _,r in pending.iterrows()]
        choice=st.selectbox("Review Emergency Item",labels,key="emergency_review_choice")
        r=pending.iloc[labels.index(choice)]
        approved=st.number_input("Approved Additional Qty",min_value=0.0,value=float(r.requested_qty),step=1.0,key=f"emer_approve_qty_{int(r.id)}")
        note=st.text_input("Review note (optional)",key=f"emer_review_note_{int(r.id)}")
        a,b=st.columns(2)
        if a.button("APPROVE / PARTIALLY APPROVE",type="primary",use_container_width=True,key=f"emer_approve_{int(r.id)}"):
            status='Approved' if float(approved)==float(r.requested_qty) else 'Partially Approved'
            with db_conn() as c: c.execute("UPDATE emergency_orders SET approved_qty=?,status=?,reviewed_by=?,reviewed_at=?,review_note=? WHERE id=?",(float(approved),status,current_user,manila_now().isoformat(timespec='seconds'),note,int(r.id)))
            audit("EMERGENCY_ORDER_REVIEWED",f"{r.order_no}; {r.branch}; {r['item']}; {status}; approved={approved}"); st.success("Emergency order reviewed and moved to History."); st.rerun()
        if b.button("REJECT",use_container_width=True,key=f"emer_reject_{int(r.id)}"):
            with db_conn() as c: c.execute("UPDATE emergency_orders SET approved_qty=0,status='Rejected',reviewed_by=?,reviewed_at=?,review_note=? WHERE id=?",(current_user,manila_now().isoformat(timespec='seconds'),note,int(r.id)))
            audit("EMERGENCY_ORDER_REJECTED",f"{r.order_no}; {r.branch}; {r['item']}"); st.success("Emergency order rejected and moved to History."); st.rerun()
    with st.expander("Emergency Order History / Delete Transaction",expanded=False):
        hist=q[(q.status!='Pending Approval') & (q.archived==0)]
        if hist.empty: st.info("No reviewed emergency transactions.")
        else:
            st.dataframe(hist[['order_no','branch','category','item','requested_qty','approved_qty','status','reviewed_by','reviewed_at']],use_container_width=True,hide_index=True)
            opts=[f"#{int(r.id)} — {r.order_no} — {r.branch} — {r['item']} — {r.status}" for _,r in hist.iterrows()]
            pick=st.selectbox("Transaction to remove from active history",opts,key="emer_archive_pick")
            rr=hist.iloc[opts.index(pick)]
            reason=st.text_input("Delete / archive reason *",key=f"emer_archive_reason_{int(rr.id)}")
            if st.button("DELETE TRANSACTION",key=f"emer_archive_{int(rr.id)}"):
                if not reason.strip(): st.error("Reason is required.")
                else:
                    with db_conn() as c: c.execute("UPDATE emergency_orders SET archived=1,archived_by=?,archived_at=?,archive_reason=? WHERE id=?",(current_user,manila_now().isoformat(timespec='seconds'),reason.strip(),int(rr.id)))
                    audit("EMERGENCY_ORDER_ARCHIVED",f"ID={int(rr.id)}; {rr.order_no}; reason={reason.strip()}"); st.success("Transaction removed from the visible history. Audit record preserved."); st.rerun()

def _stock_request_label(r):
    return f"{r.get('request_no','')} — {r.get('category','')}"


def stock_count_admin():
    st.markdown("## Stock Count Requests")
    st.caption("Request actual physical stock counts from selected branches and selected items only.")

    tab_new, tab_monitor, tab_history = st.tabs(["Create Request", "Monitor & Consolidated View", "History"])

    with tab_new:
        # Category stays outside the form so Streamlit reruns immediately when it changes.
        # This prevents items from the previous category from remaining in the selector.
        category = st.selectbox("Category *", ORDER_CATEGORIES, key="stock_count_category")
        items = category_items(category)
        with st.form(f"new_stock_count_request_{category}"):
            request_no = st.text_input("Stock Count Request No. *", placeholder="e.g. SC-2026-001").strip()
            selected_items = st.multiselect("Items to Count *", items, default=[], key=f"stock_count_items_{category}")
            branches = available_branches()
            selected_branches = st.multiselect("Branches *", branches, default=[])
            deadline_date = st.date_input("Deadline Date", value=date.today())
            deadline_time = st.time_input("Deadline Time")
            submitted = st.form_submit_button("SEND STOCK COUNT REQUEST", type="primary", use_container_width=True)
        if submitted:
            if not request_no:
                st.error("Stock Count Request No. is required.")
            elif not selected_items:
                st.error("Select at least one item to count.")
            elif not selected_branches:
                st.error("Select at least one branch.")
            else:
                try:
                    deadline = datetime.combine(deadline_date, deadline_time).isoformat()
                    created = _sb_insert("stock_count_requests", [{
                        "request_no": request_no,
                        "category": category,
                        "deadline": deadline,
                        "status": "Open",
                        "created_by": current_user,
                    }])
                    rid = created[0]["id"]
                    _sb_insert("stock_count_request_items", [
                        {"request_id": rid, "item_name": item, "item_order": i}
                        for i, item in enumerate(selected_items, start=1)
                    ])
                    _sb_insert("stock_count_request_branches", [
                        {"request_id": rid, "branch": b, "status": "Pending"}
                        for b in selected_branches
                    ])
                    audit("STOCK_COUNT_REQUEST_SENT", f"{request_no}; {category}; {', '.join(selected_branches)}")
                    st.success(f"Stock Count Request {request_no} sent to {len(selected_branches)} branch(es).")
                    st.rerun()
                except Exception as e:
                    st.error(str(e))

    def render_monitor(history=False):
        try:
            requests = _sb_request("stock_count_requests", "GET", params={"select":"*", "order":"id.desc"})
        except Exception as e:
            st.error(str(e)); return
        if history:
            requests = [r for r in requests if r.get("status") != "Open"]
        if not requests:
            st.info("No stock count requests found."); return
        labels = [_stock_request_label(r) for r in requests]
        chosen = st.selectbox("Select Stock Count Request", labels, key=f"stock_req_select_{'hist' if history else 'mon'}")
        req = requests[labels.index(chosen)]
        rid = req["id"]
        st.write(f"**Request No.:** {req.get('request_no')}  |  **Category:** {req.get('category')}  |  **Deadline:** {req.get('deadline') or '—'}")
        try:
            br = _sb_get("stock_count_request_branches", request_id=rid)
            its = _sb_get("stock_count_request_items", request_id=rid)
            entries = _sb_get("stock_count_entries", request_id=rid)
        except Exception as e:
            st.error(str(e)); return
        br_df = pd.DataFrame(br)
        if not br_df.empty:
            now = datetime.now()
            deadline = pd.to_datetime(req.get("deadline"), errors="coerce")
            def visible_status(row):
                status = row.get("status", "Pending")
                if status == "Pending" and pd.notna(deadline) and now > deadline.to_pydatetime().replace(tzinfo=None):
                    return "Late"
                return status
            br_df["Display Status"] = br_df.apply(visible_status, axis=1)
            st.markdown("### Submission Status")
            st.dataframe(br_df[["branch","Display Status","submitted_by","submitted_at"]].rename(columns={"branch":"Branch","submitted_by":"Submitted By","submitted_at":"Submitted At"}), use_container_width=True, hide_index=True)
        item_order = [x["item_name"] for x in sorted(its, key=lambda x: x.get("item_order",0))]
        branch_order = [x["branch"] for x in br]
        if entries:
            edf = pd.DataFrame(entries)
            pivot = edf.pivot_table(index="item_name", columns="branch", values="actual_stock", aggfunc="last")
            pivot = pivot.reindex(index=item_order, columns=branch_order)
        else:
            pivot = pd.DataFrame(index=item_order, columns=branch_order, dtype=float)
        pivot.index.name = "ITEM"
        numeric = pivot.apply(pd.to_numeric, errors="coerce")
        pivot["TOTAL"] = numeric.sum(axis=1, min_count=1)
        st.markdown("### Consolidated Stock Count")
        st.dataframe(pivot, use_container_width=True)
        out = io.BytesIO()
        with pd.ExcelWriter(out, engine="openpyxl") as writer:
            pivot.to_excel(writer, sheet_name="Consolidated Stock Count")
        st.download_button("⬇️ DOWNLOAD CONSOLIDATED EXCEL", out.getvalue(), f"{safe_filename(req.get('request_no'))}_Stock_Count.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
        if not history and req.get("status") == "Open":
            if st.button("CLOSE STOCK COUNT REQUEST", key=f"close_stock_{rid}"):
                try:
                    _sb_patch("stock_count_requests", {"id":rid}, {"status":"Closed", "closed_at":datetime.now().isoformat()})
                    audit("STOCK_COUNT_REQUEST_CLOSED", req.get("request_no",""))
                    st.success("Stock Count Request closed."); st.rerun()
                except Exception as e: st.error(str(e))

    with tab_monitor:
        render_monitor(False)
    with tab_history:
        render_monitor(True)


def stock_count_staff():
    branch = branch_for_user()
    st.markdown("## Stock Count")
    st.caption(f"Branch: {branch} — Enter the actual physical quantity for every requested item. Enter 0 if none.")
    try:
        branch_rows = _sb_request("stock_count_request_branches", "GET", params={"select":"*", "branch":f"eq.{branch}", "order":"id.desc"})
    except Exception as e:
        st.error(str(e)); return
    if not branch_rows:
        st.info("No Stock Count Requests have been sent to your branch."); return
    request_ids = [r["request_id"] for r in branch_rows]
    requests = []
    for rid in request_ids:
        rows = _sb_get("stock_count_requests", id=rid)
        if rows: requests.append(rows[0])
    labels = [_stock_request_label(r) for r in requests]
    chosen = st.selectbox("Select Stock Count Request", labels)
    req = requests[labels.index(chosen)]
    rid = req["id"]
    br_rec = next(r for r in branch_rows if r["request_id"] == rid)
    st.write(f"**Request No.:** {req.get('request_no')}  |  **Category:** {req.get('category')}  |  **Deadline:** {req.get('deadline') or '—'}")
    if br_rec.get("status") == "Submitted":
        st.success(f"Submitted by {br_rec.get('submitted_by')} at {br_rec.get('submitted_at')}.")
        entries = _sb_get("stock_count_entries", request_id=rid, branch=branch)
        if entries:
            st.dataframe(pd.DataFrame(entries)[["item_name","actual_stock"]].rename(columns={"item_name":"Item","actual_stock":"Actual Stock"}), use_container_width=True, hide_index=True)
        return
    if req.get("status") != "Open":
        st.warning("This Stock Count Request is already closed."); return
    items = _sb_get("stock_count_request_items", request_id=rid)
    items = sorted(items, key=lambda x: x.get("item_order",0))
    existing = _sb_get("stock_count_entries", request_id=rid, branch=branch)
    existing_map = {x["item_name"]: x.get("actual_stock") for x in existing}
    frame = pd.DataFrame({"Item":[x["item_name"] for x in items], "Actual Stock":[existing_map.get(x["item_name"], None) for x in items]})
    edited = st.data_editor(frame, hide_index=True, use_container_width=True, disabled=["Item"], column_config={"Actual Stock":st.column_config.NumberColumn("Actual Stock", min_value=0.0, step=1.0, required=True)}, key=f"stock_count_editor_{rid}_{branch}")
    st.caption("Blank = not counted yet. 0 = physically no stock.")
    if st.button("REVIEW & SUBMIT STOCK COUNT", type="primary", use_container_width=True):
        values = pd.to_numeric(edited["Actual Stock"], errors="coerce")
        if values.isna().any():
            st.error("Stock count incomplete. Enter a quantity for every requested item. Enter 0 if there is no stock.")
            return
        try:
            now = datetime.now().isoformat(timespec="seconds")
            # Insert each requested item once. Submitted requests are locked afterward.
            rows = [{"request_id":rid, "branch":branch, "item_name":str(row["Item"]), "actual_stock":float(row["Actual Stock"]), "submitted_by":current_user, "submitted_at":now} for _, row in edited.iterrows()]
            if existing:
                st.error("A stock count has already been saved for this request. Please contact Super Admin.")
                return
            _sb_insert("stock_count_entries", rows)
            _sb_patch("stock_count_request_branches", {"request_id":rid, "branch":branch}, {"status":"Submitted", "submitted_by":current_user, "submitted_at":now})
            audit("STOCK_COUNT_SUBMITTED", f"{req.get('request_no')}; {branch}")
            st.success("Stock count submitted successfully.")
            st.rerun()
        except Exception as e:
            st.error(str(e))



KAPE_OPENERS = [
    "Kape muna", "One sip at a time", "Good vibes and good coffee", "Brew, breathe, begin", "Fresh cup, fresh start",
    "Kape check", "Coffee first", "Powered by kape", "Sip, smile, work", "Mainit na kape, malinaw na isip"
]
KAPE_MIDDLES = [
    "then double-check the numbers", "at i-check ang quantity bago mag-submit", "because accuracy is teamwork", "then focus sa next task",
    "at huwag hulaan ang inventory", "because small details matter", "then make every input count", "at report agad kapag may variance",
    "because clear orders mean less hassle", "at tandaan: tama muna bago mabilis"
]
KAPE_ENDINGS = [
    "☕", "🤎", "☕🤎", "😌☕", "✨☕", "— The Perfect Blend.", "Good shift! ☕", "Keep brewing! 🤎", "Kaya mo 'yan. ☕", "Less hassle, more coffee. 🤎"
]
# Exactly 1,000 unique combinations: 10 x 10 x 10.
DAILY_BREW_MESSAGES = [f"{a}. {b}. {c}" for a in KAPE_OPENERS for b in KAPE_MIDDLES for c in KAPE_ENDINGS]

def daily_brew_message():
    # A fresh reminder is selected whenever the Dashboard is rendered/opened.
    return random.choice(DAILY_BREW_MESSAGES)

def render_daily_brew():
    st.markdown(f"""<div style='background:#FBF7F2;border:1px solid #E4D8CD;border-radius:16px;padding:18px 20px;margin:8px 0 18px 0'><b style='color:#3C271B'>☕ Kape Reminder of the Day</b><br><span style='color:#5A3F2D;font-size:1.05rem'>{daily_brew_message()}</span></div>""",unsafe_allow_html=True)

# V13.1 role-aware navigation
is_admin = current_role in ("Admin", "Super Admin")
if is_admin:
    admin_menu=["Dashboard","Orders","Stock Count Requests","Delivery Receipt Generator","Receiving & Variances","Deliveries","Transaction History","User Management"]
    if current_role == "Super Admin": admin_menu.extend(["Database Backup", "Settings"])
    mode=st.sidebar.radio("System Menu",admin_menu)
    if mode=="Dashboard":
        st.markdown("## Super Admin Dashboard")
        render_daily_brew()
        with db_conn() as c:
            total=c.execute("SELECT COUNT(*) FROM order_cycles").fetchone()[0]
            pending=c.execute("SELECT COUNT(*) FROM order_cycle_branches WHERE status='Pending'").fetchone()[0]
            recv=c.execute("SELECT COUNT(*) FROM deliveries_v2 WHERE status='For Receiving'").fetchone()[0]
        a,b,c1=st.columns(3); a.metric("Order Requests",total); b.metric("Pending Branch Submissions",pending); c1.metric("For Receiving",recv)
        st.caption("Create an Order No. under Orders, select the category and branches, then send the request to staff.")
    elif mode=="Orders": super_orders()
    elif mode=="Stock Count Requests": stock_count_admin()
    elif mode=="Delivery Receipt Generator": render_dr_generator()
    elif mode=="Receiving & Variances": receiving_variances_admin()
    elif mode=="Deliveries": deliveries_history()
    elif mode=="Transaction History": transaction_history()
    elif mode=="User Management": admin_panel()
    elif mode=="Database Backup": database_backup()
    elif mode=="Settings": branch_management()
else:
    mode=st.sidebar.radio("Branch Menu",["Dashboard","Order Requests","Stock Count","My Orders","Receive Delivery","Variances","History"])
    if mode=="Dashboard":
        st.markdown(f"## {branch_for_user()} Branch Dashboard")
        render_daily_brew()
        with db_conn() as c:
            n=c.execute("SELECT COUNT(*) FROM order_cycle_branches WHERE branch=? AND status='Pending'",(branch_for_user(),)).fetchone()[0]
        if n: st.warning(f"You have {n} new order request(s) waiting for your input.")
        st.write("Use Order Requests to enter quantities only when Super Admin sends an order request to your branch.")
    elif mode=="Order Requests": staff_orders()
    elif mode=="Stock Count": stock_count_staff()
    elif mode=="My Orders": my_orders()
    elif mode=="Receive Delivery": receive_delivery()
    elif mode=="Variances": transaction_history()
    elif mode=="History": transaction_history()

